import traceback
import uuid
from datetime import date, time, datetime
from django.conf import settings

from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.db.models import Prefetch, Q
from django.http import HttpResponse
from io import BytesIO
from openpyxl import load_workbook, Workbook

from django.db import transaction
from openpyxl.styles import Alignment, PatternFill, Font
from openpyxl.worksheet.datavalidation import DataValidation
from rest_framework.permissions import IsAuthenticated
from rest_framework.views import APIView

from apps.school.models.school import Branch, AcademicYear, Grade, Subject, Staff, Section
from apps.timetable.models import TimetableTerm, TimetablePeriod, Room, Timetable, TimetableEntry, TimetableSubstitution
from shared.mixins import CustomResponse, CustomPageNumberPagination
from shared.permissions import HasPermission
from shared.utils.logger import application_logger

from openpyxl import load_workbook

from django.db import transaction, IntegrityError

from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser


from apps.timetable.models import (
    Timetable,
    TimetableEntry,
    TimetablePeriod,
)

from shared.permissions import HasPermission
from shared.utils.logger import application_logger


class TimetableTermCreateAPIView(APIView):
    permission_classes = [
        IsAuthenticated,
        HasPermission,
    ]

    required_permission = "timetable_term.create"

    def post(self, request):
        try:
            school = request.school

            if not school:
                return CustomResponse.errorResponse(
                    description="School is required."
                )

            name = request.data.get("name")
            academic_year_id = request.data.get("academic_year_id")
            branch_id = request.data.get("branch_id")

            start_date = request.data.get("start_date")
            end_date = request.data.get("end_date")

            # ---------------------------------
            # Required fields
            # ---------------------------------

            if not name:
                return CustomResponse.errorResponse(
                    description="Term name is required."
                )

            if not academic_year_id:
                return CustomResponse.errorResponse(
                    description="Academic year is required."
                )

            if not start_date:
                return CustomResponse.errorResponse(
                    description="Start date is required."
                )

            if not end_date:
                return CustomResponse.errorResponse(
                    description="End date is required."
                )

            # ---------------------------------
            # Validate UUIDs
            # ---------------------------------

            try:
                academic_year_uuid = uuid.UUID(
                    str(academic_year_id)
                )
            except (ValueError, AttributeError):
                return CustomResponse.errorResponse(
                    description="Invalid academic_year_id."
                )

            branch = None

            if branch_id:
                try:
                    branch_uuid = uuid.UUID(str(branch_id))
                except (ValueError, AttributeError):
                    return CustomResponse.errorResponse(
                        description="Invalid branch_id."
                    )

                branch = Branch.objects.filter(
                    id=branch_uuid,
                    school=school,
                ).first()

                if not branch:
                    return CustomResponse.errorResponse(
                        description="Invalid branch."
                    )

            # ---------------------------------
            # Validate dates
            # ---------------------------------

            if start_date > end_date:
                return CustomResponse.errorResponse(
                    description="Start date cannot be after end date."
                )

            # ---------------------------------
            # Academic Year
            # ---------------------------------

            academic_year = AcademicYear.objects.filter(
                id=academic_year_uuid,
                school=school,
            ).first()

            if not academic_year:
                return CustomResponse.errorResponse(
                    description="Invalid academic year."
                )

            # ---------------------------------
            # Duplicate term
            # ---------------------------------

            if TimetableTerm.objects.filter(
                school=school,
                academic_year=academic_year,
                name__iexact=name.strip(),
            ).exists():
                return CustomResponse.errorResponse(
                    description=(
                        "Timetable term with this name "
                        "already exists."
                    )
                )

            # ---------------------------------
            # Create
            # ---------------------------------

            term = TimetableTerm.objects.create(
                school=school,
                branch=branch,
                academic_year=academic_year,
                name=name.strip(),
                start_date=start_date,
                end_date=end_date,
                status=TimetableTerm.Status.ACTIVE,
            )

            application_logger.info(
                "timetable_term_created",
                school_id=str(school.id),
                term_id=str(term.id),
                academic_year_id=str(
                    academic_year.id
                ),
            )

            return CustomResponse.successResponse(
                data={
                    "id": term.id,
                    "name": term.name,
                    "academic_year_id": term.academic_year_id,
                    "branch_id": term.branch_id,
                    "start_date": term.start_date,
                    "end_date": term.end_date,
                    "status": term.status,
                },
                description="Timetable term created successfully.",
            )

        except Exception as e:

            application_logger.exception(
                "timetable_term_create_failed",
                error=str(e),
                school_id=(
                    str(request.school.id)
                    if getattr(request, "school", None)
                    else None
                ),
            )

            return CustomResponse.errorResponse(
                description="Failed to create timetable term."
            )

class TimetableTermListAPIView(APIView):
    permission_classes = [
        IsAuthenticated,
        HasPermission,
    ]

    required_permission = "timetable_term.view"

    def get(self, request):
        try:
            school = request.school

            academic_year_id = request.query_params.get(
                "academic_year_id"
            )
            branch_id = request.query_params.get("branch_id")
            status = request.query_params.get("status")
            search = request.query_params.get("search")

            queryset = (
                TimetableTerm.objects
                .filter(
                    school=school
                )
                .select_related(
                    "academic_year",
                    "branch",
                )
                .order_by(
                    "-academic_year__start_date",
                    "start_date",
                )
            )

            # ---------------------------------
            # Academic year
            # ---------------------------------

            if academic_year_id:
                try:
                    academic_year_uuid = uuid.UUID(
                        str(academic_year_id)
                    )
                except (ValueError, AttributeError):
                    return CustomResponse.errorResponse(
                        description="Invalid academic_year_id."
                    )

                queryset = queryset.filter(
                    academic_year_id=academic_year_uuid
                )

            # ---------------------------------
            # Branch
            # ---------------------------------

            if branch_id:
                try:
                    branch_uuid = uuid.UUID(
                        str(branch_id)
                    )
                except (ValueError, AttributeError):
                    return CustomResponse.errorResponse(
                        description="Invalid branch_id."
                    )

                queryset = queryset.filter(
                    branch_id=branch_uuid
                )

            # ---------------------------------
            # Status
            # ---------------------------------

            if status:
                valid_statuses = [
                    choice[0]
                    for choice in TimetableTerm.Status.choices
                ]

                if status not in valid_statuses:
                    return CustomResponse.errorResponse(
                        description="Invalid status."
                    )

                queryset = queryset.filter(
                    status=status
                )

            # ---------------------------------
            # Search
            # ---------------------------------

            if search:
                queryset = queryset.filter(
                    name__icontains=search.strip()
                )

            total = queryset.count()

            # ---------------------------------
            # Pagination
            # ---------------------------------

            paginator = CustomPageNumberPagination()

            page = paginator.paginate_queryset(
                queryset,
                request,
            )

            data = []

            for term in page:
                data.append({
                    "id": term.id,
                    "name": term.name,

                    "academic_year": {
                        "id": term.academic_year.id,
                        "name": term.academic_year.name,
                    },

                    "branch": (
                        {
                            "id": term.branch.id,
                            "name": term.branch.name,
                        }
                        if term.branch
                        else None
                    ),

                    "start_date": term.start_date,
                    "end_date": term.end_date,
                    "status": term.status,
                })

            application_logger.info(
                "timetable_term_list_fetched",
                school_id=str(school.id),
                academic_year_id=academic_year_id,
                branch_id=branch_id,
                status=status,
                total=total,
            )

            return CustomResponse.successResponse(
                data=data,
                total=total,
                description="Timetable terms fetched successfully.",
            )

        except Exception as e:

            application_logger.exception(
                "timetable_term_list_failed",
                error=str(e),
                school_id=(
                    str(school.id)
                    if school
                    else None
                ),
            )

            return CustomResponse.errorResponse(
                description="Failed to fetch timetable terms."
            )

class TimetableTermUpdateAPIView(APIView):
    permission_classes = [
        IsAuthenticated,
        HasPermission,
    ]

    required_permission = "timetable_term.update"

    def put(self, request,term_id):
        try:
            school = request.school


            if not term_id:
                return CustomResponse.errorResponse(
                    description="term_id is required."
                )

            try:
                term_uuid = uuid.UUID(str(term_id))
            except (ValueError, AttributeError):
                return CustomResponse.errorResponse(
                    description="Invalid term_id."
                )

            term = TimetableTerm.objects.filter(
                id=term_uuid,
                school=school,
            ).first()

            if not term:
                return CustomResponse.errorResponse(
                    description="Timetable term not found."
                )

            name = request.data.get("name")
            academic_year_id = request.data.get(
                "academic_year_id"
            )
            branch_id = request.data.get("branch_id")

            start_date = request.data.get("start_date")
            end_date = request.data.get("end_date")
            status = request.data.get("status")

            # ---------------------------------
            # Name
            # ---------------------------------

            if name is not None:
                name = name.strip()

                if not name:
                    return CustomResponse.errorResponse(
                        description="Term name cannot be empty."
                    )

                if TimetableTerm.objects.filter(
                    school=school,
                    academic_year=term.academic_year,
                    name__iexact=name,
                ).exclude(
                    id=term.id
                ).exists():
                    return CustomResponse.errorResponse(
                        description=(
                            "Timetable term with this name "
                            "already exists."
                        )
                    )

                term.name = name

            # ---------------------------------
            # Academic Year
            # ---------------------------------

            if academic_year_id:

                try:
                    academic_year_uuid = uuid.UUID(
                        str(academic_year_id)
                    )
                except (ValueError, AttributeError):
                    return CustomResponse.errorResponse(
                        description="Invalid academic_year_id."
                    )

                academic_year = AcademicYear.objects.filter(
                    id=academic_year_uuid,
                    school=school,
                ).first()

                if not academic_year:
                    return CustomResponse.errorResponse(
                        description="Invalid academic year."
                    )

                if TimetableTerm.objects.filter(
                    school=school,
                    academic_year=academic_year,
                    name__iexact=term.name,
                ).exclude(
                    id=term.id
                ).exists():
                    return CustomResponse.errorResponse(
                        description=(
                            "Timetable term with this name "
                            "already exists for this academic year."
                        )
                    )

                term.academic_year = academic_year

            # ---------------------------------
            # Branch
            # ---------------------------------

            if branch_id is not None:

                if branch_id:
                    try:
                        branch_uuid = uuid.UUID(
                            str(branch_id)
                        )
                    except (ValueError, AttributeError):
                        return CustomResponse.errorResponse(
                            description="Invalid branch_id."
                        )

                    branch = Branch.objects.filter(
                        id=branch_uuid,
                        school=school,
                    ).first()

                    if not branch:
                        return CustomResponse.errorResponse(
                            description="Invalid branch."
                        )

                    term.branch = branch

                else:
                    term.branch = None

            # ---------------------------------
            # Dates
            # ---------------------------------

            final_start_date = (
                start_date
                if start_date is not None
                else term.start_date
            )

            final_end_date = (
                end_date
                if end_date is not None
                else term.end_date
            )

            if final_start_date > final_end_date:
                return CustomResponse.errorResponse(
                    description=(
                        "Start date cannot be after end date."
                    )
                )

            if start_date is not None:
                term.start_date = start_date

            if end_date is not None:
                term.end_date = end_date

            # ---------------------------------
            # Status
            # ---------------------------------

            if status is not None:

                valid_statuses = [
                    choice[0]
                    for choice in TimetableTerm.Status.choices
                ]

                if status not in valid_statuses:
                    return CustomResponse.errorResponse(
                        description="Invalid status."
                    )

                term.status = status

            term.save()

            application_logger.info(
                "timetable_term_updated",
                school_id=str(school.id),
                term_id=str(term.id),
            )

            return CustomResponse.successResponse(
                data={
                    "id": term.id,
                    "name": term.name,
                    "academic_year_id": (
                        term.academic_year_id
                    ),
                    "branch_id": term.branch_id,
                    "start_date": term.start_date,
                    "end_date": term.end_date,
                    "status": term.status,
                },
                description="Timetable term updated successfully.",
            )

        except Exception as e:

            application_logger.exception(
                "timetable_term_update_failed",
                error=str(e),
                school_id=(
                    str(school.id)
                    if school
                    else None
                ),
            )

            return CustomResponse.errorResponse(
                description="Failed to update timetable term."
            )


class TimetablePeriodCreateAPIView(APIView):
    permission_classes = [
        IsAuthenticated,
        HasPermission,
    ]

    required_permission = "timetable_period.create"

    def post(self, request):
        try:
            school = request.school

            if not school:
                return CustomResponse.errorResponse(
                    description="School is required."
                )

            name = request.data.get("name")
            period_number = request.data.get("period_number")
            start_time = request.data.get("start_time")
            end_time = request.data.get("end_time")
            is_break = request.data.get("is_break", False)

            # ---------------------------------
            # Required fields
            # ---------------------------------

            if not name:
                return CustomResponse.errorResponse(
                    description="Period name is required."
                )

            if period_number is None:
                return CustomResponse.errorResponse(
                    description="Period number is required."
                )

            if not start_time:
                return CustomResponse.errorResponse(
                    description="Start time is required."
                )

            if not end_time:
                return CustomResponse.errorResponse(
                    description="End time is required."
                )

            # ---------------------------------
            # Period number
            # ---------------------------------

            try:
                period_number = int(period_number)
            except (ValueError, TypeError):
                return CustomResponse.errorResponse(
                    description="Invalid period number."
                )

            if period_number <= 0:
                return CustomResponse.errorResponse(
                    description="Period number must be greater than 0."
                )

            # ---------------------------------
            # Validate time
            # ---------------------------------

            if start_time >= end_time:
                return CustomResponse.errorResponse(
                    description="Start time must be before end time."
                )

            # ---------------------------------
            # Duplicate period number
            # ---------------------------------

            if TimetablePeriod.objects.filter(
                school=school,
                period_number=period_number,
            ).exists():
                return CustomResponse.errorResponse(
                    description=(
                        "A period with this period number "
                        "already exists."
                    )
                )

            # ---------------------------------
            # Overlapping period
            # ---------------------------------

            overlapping_period = (
                TimetablePeriod.objects
                .filter(
                    school=school,
                    start_time__lt=end_time,
                    end_time__gt=start_time,
                )
                .exists()
            )

            if overlapping_period:
                return CustomResponse.errorResponse(
                    description=(
                        "This period overlaps with an "
                        "existing period."
                    )
                )

            # ---------------------------------
            # Validate break
            # ---------------------------------

            if not isinstance(is_break, bool):
                is_break = str(is_break).lower() == "true"

            # ---------------------------------
            # Create
            # ---------------------------------

            period = TimetablePeriod.objects.create(
                school=school,
                name=name.strip(),
                period_number=period_number,
                start_time=start_time,
                end_time=end_time,
                is_break=is_break,
                status=TimetablePeriod.Status.ACTIVE,
            )

            application_logger.info(
                "timetable_period_created",
                school_id=str(school.id),
                period_id=str(period.id),
            )

            return CustomResponse.successResponse(
                data={
                    "id": period.id,
                    "name": period.name,
                    "period_number": period.period_number,
                    "start_time": period.start_time,
                    "end_time": period.end_time,
                    "is_break": period.is_break,
                    "status": period.status,
                },
                description="Timetable period created successfully.",
            )

        except Exception as e:
            application_logger.exception(
                "timetable_period_create_failed",
                error=str(e),
                school_id=(
                    str(request.school.id)
                    if getattr(request, "school", None)
                    else None
                ),
            )

            return CustomResponse.errorResponse(
                description="Failed to create timetable period."
            )

class TimetablePeriodListAPIView(APIView):
    permission_classes = [
        IsAuthenticated,
        HasPermission,
    ]

    required_permission = "timetable_period.view"

    def get(self, request):
        try:
            school = request.school

            status = request.query_params.get("status")
            is_break = request.query_params.get("is_break")
            search = request.query_params.get("search")

            queryset = (
                TimetablePeriod.objects
                .filter(
                    school=school
                )
                .order_by("period_number")
            )

            # ---------------------------------
            # Status
            # ---------------------------------

            if status:
                valid_statuses = [
                    choice[0]
                    for choice in TimetablePeriod.Status.choices
                ]

                if status not in valid_statuses:
                    return CustomResponse.errorResponse(
                        description="Invalid status."
                    )

                queryset = queryset.filter(
                    status=status
                )

            # ---------------------------------
            # Break filter
            # ---------------------------------

            if is_break is not None:
                if str(is_break).lower() not in [
                    "true",
                    "false",
                ]:
                    return CustomResponse.errorResponse(
                        description="Invalid is_break value."
                    )

                queryset = queryset.filter(
                    is_break=str(is_break).lower() == "true"
                )

            # ---------------------------------
            # Search
            # ---------------------------------

            if search:
                queryset = queryset.filter(
                    name__icontains=search.strip()
                )

            total = queryset.count()

            # ---------------------------------
            # Pagination
            # ---------------------------------

            paginator = CustomPageNumberPagination()

            page = paginator.paginate_queryset(
                queryset,
                request,
            )

            data = []

            for period in page:
                data.append({
                    "id": period.id,
                    "name": period.name,
                    "period_number": period.period_number,
                    "start_time": period.start_time,
                    "end_time": period.end_time,
                    "is_break": period.is_break,
                    "status": period.status,
                })

            application_logger.info(
                "timetable_period_list_fetched",
                school_id=str(school.id),
                status=status,
                is_break=is_break,
                total=total,
            )

            return CustomResponse.successResponse(
                data=data,
                total=total,
                description="Timetable periods fetched successfully.",
            )

        except Exception as e:
            application_logger.exception(
                "timetable_period_list_failed",
                error=str(e),
                school_id=(
                    str(school.id)
                    if school
                    else None
                ),
            )

            return CustomResponse.errorResponse(
                description="Failed to fetch timetable periods."
            )


class TimetablePeriodUpdateAPIView(APIView):
    permission_classes = [
        IsAuthenticated,
        HasPermission,
    ]

    required_permission = "timetable_period.update"

    def put(self, request,period_id):
        try:
            school = request.school


            if not period_id:
                return CustomResponse.errorResponse(
                    description="period_id is required."
                )

            try:
                period_uuid = uuid.UUID(str(period_id))
            except (ValueError, AttributeError):
                return CustomResponse.errorResponse(
                    description="Invalid period_id."
                )

            period = TimetablePeriod.objects.filter(
                id=period_uuid,
                school=school,
            ).first()

            if not period:
                return CustomResponse.errorResponse(
                    description="Timetable period not found."
                )

            name = request.data.get("name")
            period_number = request.data.get("period_number")
            start_time = request.data.get("start_time")
            end_time = request.data.get("end_time")
            is_break = request.data.get("is_break")
            status = request.data.get("status")

            # ---------------------------------
            # Name
            # ---------------------------------

            if name is not None:
                name = name.strip()

                if not name:
                    return CustomResponse.errorResponse(
                        description="Period name cannot be empty."
                    )

                period.name = name

            # ---------------------------------
            # Period number
            # ---------------------------------

            if period_number is not None:

                try:
                    period_number = int(period_number)
                except (ValueError, TypeError):
                    return CustomResponse.errorResponse(
                        description="Invalid period number."
                    )

                if period_number <= 0:
                    return CustomResponse.errorResponse(
                        description=(
                            "Period number must be greater than 0."
                        )
                    )

                if TimetablePeriod.objects.filter(
                    school=school,
                    period_number=period_number,
                ).exclude(
                    id=period.id
                ).exists():
                    return CustomResponse.errorResponse(
                        description=(
                            "A period with this period number "
                            "already exists."
                        )
                    )

                period.period_number = period_number

            # ---------------------------------
            # Final times
            # ---------------------------------

            final_start_time = (
                start_time
                if start_time is not None
                else period.start_time
            )

            final_end_time = (
                end_time
                if end_time is not None
                else period.end_time
            )

            if final_start_time >= final_end_time:
                return CustomResponse.errorResponse(
                    description="Start time must be before end time."
                )

            # ---------------------------------
            # Overlapping period
            # ---------------------------------

            overlapping_period = (
                TimetablePeriod.objects
                .filter(
                    school=school,
                    start_time__lt=final_end_time,
                    end_time__gt=final_start_time,
                )
                .exclude(
                    id=period.id
                )
                .exists()
            )

            if overlapping_period:
                return CustomResponse.errorResponse(
                    description=(
                        "This period overlaps with an "
                        "existing period."
                    )
                )

            if start_time is not None:
                period.start_time = start_time

            if end_time is not None:
                period.end_time = end_time

            # ---------------------------------
            # Break
            # ---------------------------------

            if is_break is not None:

                if not isinstance(is_break, bool):
                    is_break = (
                        str(is_break).lower() == "true"
                    )

                period.is_break = is_break

            # ---------------------------------
            # Status
            # ---------------------------------

            if status is not None:

                valid_statuses = [
                    choice[0]
                    for choice in TimetablePeriod.Status.choices
                ]

                if status not in valid_statuses:
                    return CustomResponse.errorResponse(
                        description="Invalid status."
                    )

                period.status = status

            period.save()

            application_logger.info(
                "timetable_period_updated",
                school_id=str(school.id),
                period_id=str(period.id),
            )

            return CustomResponse.successResponse(
                data={
                    "id": period.id,
                    "name": period.name,
                    "period_number": period.period_number,
                    "start_time": period.start_time,
                    "end_time": period.end_time,
                    "is_break": period.is_break,
                    "status": period.status,
                },
                description="Timetable period updated successfully.",
            )

        except Exception as e:
            application_logger.exception(
                "timetable_period_update_failed",
                error=str(e),
                school_id=(
                    str(school.id)
                    if school
                    else None
                ),
            )

            return CustomResponse.errorResponse(
                description="Failed to update timetable period."
            )

class TimetablePeriodDeleteAPIView(APIView):
    permission_classes = [
        IsAuthenticated,
        HasPermission,
    ]

    required_permission = "timetable_period.delete"

    def delete(self, request,period_id):
        try:
            school = request.school


            if not period_id:
                return CustomResponse.errorResponse(
                    description="period_id is required."
                )

            try:
                period_uuid = uuid.UUID(str(period_id))
            except (ValueError, AttributeError):
                return CustomResponse.errorResponse(
                    description="Invalid period_id."
                )

            period = TimetablePeriod.objects.filter(
                id=period_uuid,
                school=school,
            ).first()

            if not period:
                return CustomResponse.errorResponse(
                    description="Timetable period not found."
                )

            period.status = TimetablePeriod.Status.INACTIVE

            period.save(
                update_fields=[
                    "status",
                    "updated_at",
                ]
            )

            application_logger.info(
                "timetable_period_deactivated",
                school_id=str(school.id),
                period_id=str(period.id),
            )

            return CustomResponse.successResponse(
                data={
                    "id": period.id,
                    "status": period.status,
                },
                description="Timetable period deactivated successfully.",
            )

        except Exception as e:
            application_logger.exception(
                "timetable_period_deactivate_failed",
                error=str(e),
                school_id=(
                    str(school.id)
                    if school
                    else None
                ),
            )

            return CustomResponse.errorResponse(
                description="Failed to deactivate timetable period."
            )

class RoomCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = "room.create"

    def post(self, request):
        try:
            school = request.school

            name = request.data.get("name")
            room_type = request.data.get("room_type")
            location = request.data.get("location")
            capacity = request.data.get("capacity", 0)
            facilities = request.data.get("facilities")
            branch_id = request.data.get("branch_id")

            if not name:
                return CustomResponse.errorResponse(
                    message="Room name is required"
                )

            if not room_type:
                return CustomResponse.errorResponse(
                    message="Room type is required"
                )

            valid_room_types = dict(Room.ROOM_TYPE_CHOICES)

            if room_type not in valid_room_types:
                return CustomResponse.errorResponse(
                    message="Invalid room type"
                )

            try:
                capacity = int(capacity)
            except (TypeError, ValueError):
                return CustomResponse.errorResponse(
                    message="Capacity must be a valid number"
                )

            if capacity < 0:
                return CustomResponse.errorResponse(
                    message="Capacity cannot be negative"
                )

            branch = None

            if branch_id:
                try:
                    branch_id = uuid.UUID(str(branch_id))
                except ValueError:
                    return CustomResponse.errorResponse(
                        message="Invalid branch_id"
                    )

                branch = Branch.objects.filter(
                    id=branch_id,
                    school=school
                ).first()

                if not branch:
                    return CustomResponse.errorResponse(
                        message="Branch not found"
                    )

            if Room.objects.filter(
                school=school,
                name__iexact=name.strip()
            ).exists():
                return CustomResponse.errorResponse(
                    message="Room with this name already exists"
                )

            room = Room.objects.create(
                school=school,
                branch=branch,
                name=name.strip(),
                room_type=room_type,
                location=location,
                capacity=capacity,
                facilities=facilities,
                status="ACTIVE"
            )

            application_logger.info(
                "room_created | room_id=%s | school_id=%s",
                room.id,
                school.id
            )

            return CustomResponse.successResponse(
                message="Room created successfully",
                data={
                    "id": str(room.id),
                    "name": room.name,
                    "room_type": room.room_type,
                    "location": room.location,
                    "capacity": room.capacity,
                    "facilities": room.facilities,
                    "status": room.status,
                    "branch_id": str(room.branch_id) if room.branch_id else None
                }
            )

        except Exception:
            application_logger.exception("room_create_failed")

            return CustomResponse.errorResponse(
                message="Failed to create room"
            )

class RoomListAPIView(APIView):
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = "room.view"

    def get(self, request):
        try:
            school = request.school

            search = request.GET.get("search")
            room_type = request.GET.get("room_type")
            branch_id = request.GET.get("branch_id")
            status = request.GET.get("status")

            rooms = Room.objects.filter(
                school=school
            ).select_related("branch").order_by("name")

            if search:
                rooms = rooms.filter(
                    Q(name__icontains=search) |
                    Q(location__icontains=search) |
                    Q(facilities__icontains=search)
                )

            if room_type:
                rooms = rooms.filter(
                    room_type=room_type
                )

            if status:
                rooms = rooms.filter(
                    status=status
                )

            if branch_id:
                try:
                    branch_id = uuid.UUID(str(branch_id))
                except ValueError:
                    return CustomResponse.errorResponse(
                        message="Invalid branch_id"
                    )

                rooms = rooms.filter(
                    branch_id=branch_id
                )

            paginator = CustomPageNumberPagination()
            paginated_rooms = paginator.paginate_queryset(
                rooms,
                request
            )

            data = []

            for room in paginated_rooms:
                data.append({
                    "id": str(room.id),
                    "name": room.name,
                    "room_type": room.room_type,
                    "room_type_display": room.get_room_type_display(),
                    "location": room.location,
                    "capacity": room.capacity,
                    "facilities": room.facilities,
                    "status": room.status,
                    "branch": {
                        "id": str(room.branch_id),
                        "name": room.branch.name
                    } if room.branch else None
                })

            return paginator.get_paginated_response(data)

        except Exception:
            application_logger.exception("room_list_failed")

            return CustomResponse.errorResponse(
                message="Failed to fetch rooms"
            )

class RoomUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = "room.update"

    def put(self, request):
        try:
            school = request.school
            room_id = request.data.get("room_id")

            if not room_id:
                return CustomResponse.errorResponse(
                    message="room_id is required"
                )

            try:
                room_id = uuid.UUID(str(room_id))
            except ValueError:
                return CustomResponse.errorResponse(
                    message="Invalid room_id"
                )

            room = Room.objects.filter(
                id=room_id,
                school=school
            ).first()

            if not room:
                return CustomResponse.errorResponse(
                    message="Room not found"
                )

            name = request.data.get("name")
            room_type = request.data.get("room_type")
            location = request.data.get("location")
            capacity = request.data.get("capacity")
            facilities = request.data.get("facilities")
            branch_id = request.data.get("branch_id")
            status = request.data.get("status")

            if name is not None:
                name = name.strip()

                if not name:
                    return CustomResponse.errorResponse(
                        message="Room name cannot be empty"
                    )

                if Room.objects.filter(
                    school=school,
                    name__iexact=name
                ).exclude(id=room.id).exists():
                    return CustomResponse.errorResponse(
                        message="Room with this name already exists"
                    )

                room.name = name

            if room_type is not None:
                valid_room_types = dict(Room.ROOM_TYPE_CHOICES)

                if room_type not in valid_room_types:
                    return CustomResponse.errorResponse(
                        message="Invalid room type"
                    )

                room.room_type = room_type

            if location is not None:
                room.location = location

            if capacity is not None:
                try:
                    capacity = int(capacity)
                except (TypeError, ValueError):
                    return CustomResponse.errorResponse(
                        message="Capacity must be a valid number"
                    )

                if capacity < 0:
                    return CustomResponse.errorResponse(
                        message="Capacity cannot be negative"
                    )

                room.capacity = capacity

            if facilities is not None:
                room.facilities = facilities

            if branch_id is not None:
                if branch_id:
                    try:
                        branch_id = uuid.UUID(str(branch_id))
                    except ValueError:
                        return CustomResponse.errorResponse(
                            message="Invalid branch_id"
                        )

                    branch = Branch.objects.filter(
                        id=branch_id,
                        school=school
                    ).first()

                    if not branch:
                        return CustomResponse.errorResponse(
                            message="Branch not found"
                        )

                    room.branch = branch
                else:
                    room.branch = None

            if status is not None:
                valid_statuses = dict(Room.STATUS_CHOICES)

                if status not in valid_statuses:
                    return CustomResponse.errorResponse(
                        message="Invalid status"
                    )

                room.status = status

            room.save()

            application_logger.info(
                "room_updated | room_id=%s | school_id=%s",
                room.id,
                school.id
            )

            return CustomResponse.successResponse(
                message="Room updated successfully",
                data={
                    "id": str(room.id),
                    "name": room.name,
                    "room_type": room.room_type,
                    "location": room.location,
                    "capacity": room.capacity,
                    "facilities": room.facilities,
                    "status": room.status,
                    "branch_id": (
                        str(room.branch_id)
                        if room.branch_id else None
                    )
                }
            )

        except Exception:
            application_logger.exception("room_update_failed")

            return CustomResponse.errorResponse(
                message="Failed to update room"
            )

class RoomDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = "room.delete"

    def delete(self, request):
        try:
            school = request.school
            room_id = request.GET.get("room_id")

            if not room_id:
                return CustomResponse.errorResponse(
                    message="room_id is required"
                )

            try:
                room_id = uuid.UUID(str(room_id))
            except ValueError:
                return CustomResponse.errorResponse(
                    message="Invalid room_id"
                )

            room = Room.objects.filter(
                id=room_id,
                school=school
            ).first()

            if not room:
                return CustomResponse.errorResponse(
                    message="Room not found"
                )

            room.status = "INACTIVE"
            room.save(update_fields=["status"])

            application_logger.info(
                "room_deactivated | room_id=%s | school_id=%s",
                room.id,
                school.id
            )

            return CustomResponse.successResponse(
                message="Room deactivated successfully"
            )

        except Exception:
            application_logger.exception("room_delete_failed")

            return CustomResponse.errorResponse(
                message="Failed to deactivate room"
            )

class RoomExcelImportAPIView(APIView):
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = "room.create"

    def post(self, request):
        try:
            school = request.school
            excel_file = request.FILES.get("file")

            if not excel_file:
                return CustomResponse.errorResponse(
                    message="Excel file is required"
                )

            if not excel_file.name.lower().endswith(
                (".xlsx", ".xls")
            ):
                return CustomResponse.errorResponse(
                    message="Only Excel files are allowed"
                )

            try:
                workbook = load_workbook(
                    excel_file,
                    data_only=True
                )
            except Exception:
                return CustomResponse.errorResponse(
                    message="Invalid Excel file"
                )

            if "Rooms" not in workbook.sheetnames:
                return CustomResponse.errorResponse(
                    message="Rooms sheet is required"
                )

            sheet = workbook["Rooms"]

            rows = list(sheet.iter_rows(values_only=True))

            if not rows:
                return CustomResponse.errorResponse(
                    message="Excel file is empty"
                )

            headers = [
                str(value).strip().lower()
                if value is not None else ""
                for value in rows[0]
            ]

            required_columns = [
                "name",
                "room_type",
                "capacity",
            ]

            missing_columns = [
                column
                for column in required_columns
                if column not in headers
            ]

            if missing_columns:
                return CustomResponse.errorResponse(
                    message="Required columns are missing",
                    data={
                        "missing_columns": missing_columns
                    }
                )

            column_index = {
                column: headers.index(column)
                for column in headers
                if column
            }

            valid_room_types = dict(Room.ROOM_TYPE_CHOICES)
            valid_statuses = dict(Room.STATUS_CHOICES)

            errors = []
            valid_rows = []

            for row_number, row in enumerate(
                rows[1:],
                start=2
            ):
                # Skip completely empty rows
                if not any(
                    value is not None
                    for value in row
                ):
                    continue

                def get_value(column):
                    index = column_index.get(column)

                    if index is None or index >= len(row):
                        return None

                    value = row[index]

                    if isinstance(value, str):
                        return value.strip()

                    return value

                name = get_value("name")
                room_type = get_value("room_type")
                location = get_value("location")
                capacity = get_value("capacity")
                facilities = get_value("facilities")
                status = get_value("status") or "ACTIVE"
                branch_id = get_value("branch_id")

                row_errors = []

                # -------------------------
                # Name validation
                # -------------------------

                if not name:
                    row_errors.append(
                        "Room name is required"
                    )

                # -------------------------
                # Room type validation
                # -------------------------

                if not room_type:
                    row_errors.append(
                        "Room type is required"
                    )
                elif room_type not in valid_room_types:
                    row_errors.append(
                        f"Invalid room_type '{room_type}'. "
                        f"Allowed values: "
                        f"{', '.join(valid_room_types.keys())}"
                    )

                # -------------------------
                # Capacity validation
                # -------------------------

                try:
                    capacity = int(capacity)
                    if capacity < 0:
                        row_errors.append(
                            "Capacity cannot be negative"
                        )
                except (TypeError, ValueError):
                    row_errors.append(
                        "Capacity must be a valid number"
                    )

                # -------------------------
                # Status validation
                # -------------------------

                if status not in valid_statuses:
                    row_errors.append(
                        f"Invalid status '{status}'"
                    )

                # -------------------------
                # Branch validation
                # -------------------------

                branch = None

                if branch_id:

                    try:
                        branch_uuid = uuid.UUID(
                            str(branch_id)
                        )
                    except ValueError:
                        row_errors.append(
                            "Invalid branch_id"
                        )
                        branch_uuid = None

                    if branch_uuid:
                        branch = Branch.objects.filter(
                            id=branch_uuid,
                            school=school
                        ).first()

                        if not branch:
                            row_errors.append(
                                "Branch not found"
                            )

                # -------------------------
                # Duplicate in database
                # -------------------------

                if name:
                    if Room.objects.filter(
                        school=school,
                        name__iexact=name
                    ).exists():
                        row_errors.append(
                            f"Room '{name}' already exists"
                        )

                # -------------------------
                # Duplicate inside Excel
                # -------------------------

                if name:
                    duplicate_excel = any(
                        item["name"].lower() == name.lower()
                        for item in valid_rows
                    )

                    if duplicate_excel:
                        row_errors.append(
                            f"Duplicate room '{name}' "
                            f"found in Excel"
                        )

                if row_errors:
                    errors.append({
                        "row": row_number,
                        "name": name,
                        "errors": row_errors
                    })
                    continue

                valid_rows.append({
                    "name": name,
                    "room_type": room_type,
                    "location": location,
                    "capacity": capacity,
                    "facilities": facilities,
                    "status": status,
                    "branch": branch
                })

            # --------------------------------
            # Do not insert anything if errors
            # --------------------------------

            if errors:
                return CustomResponse.errorResponse(
                    message="Excel validation failed",
                    data={
                        "total_rows": len(rows) - 1,
                        "valid_rows": len(valid_rows),
                        "error_rows": len(errors),
                        "errors": errors
                    }
                )

            # --------------------------------
            # Create rooms atomically
            # --------------------------------

            with transaction.atomic():

                rooms = Room.objects.bulk_create([
                    Room(
                        school=school,
                        branch=item["branch"],
                        name=item["name"],
                        room_type=item["room_type"],
                        location=item["location"],
                        capacity=item["capacity"],
                        facilities=item["facilities"],
                        status=item["status"]
                    )
                    for item in valid_rows
                ])

            application_logger.info(
                "rooms_excel_imported | school_id=%s | count=%s",
                school.id,
                len(rooms)
            )

            return CustomResponse.successResponse(
                message="Rooms imported successfully",
                data={
                    "total_rooms": len(rooms),
                    "rooms": [
                        {
                            "id": str(room.id),
                            "name": room.name,
                            "room_type": room.room_type,
                            "location": room.location,
                            "capacity": room.capacity,
                            "facilities": room.facilities,
                            "status": room.status,
                            "branch_id": (
                                str(room.branch_id)
                                if room.branch_id
                                else None
                            )
                        }
                        for room in rooms
                    ]
                }
            )

        except Exception:
            application_logger.exception(
                "rooms_excel_import_failed"
            )

            return CustomResponse.errorResponse(
                message="Failed to import rooms"
            )

class RoomExcelTemplateDownloadAPIView(APIView):
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = "room.create"

    def get(self, request):
        try:
            workbook = Workbook()

            sheet = workbook.active
            sheet.title = "Rooms"

            # -------------------------
            # Headers
            # -------------------------

            headers = [
                "name",
                "room_type",
                "location",
                "capacity",
                "facilities",
                "status",
                "branch_id",
            ]

            for column, header in enumerate(headers, start=1):
                cell = sheet.cell(
                    row=1,
                    column=column,
                    value=header
                )

                cell.font = Font(bold=True)
                cell.fill = PatternFill(
                    start_color="D9EAF7",
                    end_color="D9EAF7",
                    fill_type="solid"
                )
                cell.alignment = Alignment(
                    horizontal="center"
                )

            # -------------------------
            # Example rows
            # -------------------------

            example_rows = [
                [
                    "Room 101",
                    "CLASSROOM",
                    "First Floor",
                    40,
                    "Projector",
                    "ACTIVE",
                    ""
                ],
                [
                    "Science Lab",
                    "LAB",
                    "Ground Floor",
                    30,
                    "Computers, Projector",
                    "ACTIVE",
                    ""
                ],
            ]

            for row in example_rows:
                sheet.append(row)

            # -------------------------
            # Column widths
            # -------------------------

            sheet.column_dimensions["A"].width = 25
            sheet.column_dimensions["B"].width = 20
            sheet.column_dimensions["C"].width = 25
            sheet.column_dimensions["D"].width = 12
            sheet.column_dimensions["E"].width = 35
            sheet.column_dimensions["F"].width = 15
            sheet.column_dimensions["G"].width = 40

            # -------------------------
            # Room type dropdown
            # -------------------------

            room_types = [
                choice[0]
                for choice in Room.ROOM_TYPE_CHOICES
            ]

            room_type_formula = (
                '"' + ",".join(room_types) + '"'
            )

            room_type_validation = DataValidation(
                type="list",
                formula1=room_type_formula,
                allow_blank=False
            )

            room_type_validation.error = (
                "Please select a valid room type."
            )

            room_type_validation.errorTitle = (
                "Invalid Room Type"
            )

            sheet.add_data_validation(
                room_type_validation
            )

            room_type_validation.add(
                "B2:B1000"
            )

            # -------------------------
            # Status dropdown
            # -------------------------

            statuses = [
                choice[0]
                for choice in Room.RoomType
            ]

            status_formula = (
                '"' + ",".join(statuses) + '"'
            )

            status_validation = DataValidation(
                type="list",
                formula1=status_formula,
                allow_blank=False
            )

            status_validation.error = (
                "Please select a valid status."
            )

            status_validation.errorTitle = (
                "Invalid Status"
            )

            sheet.add_data_validation(
                status_validation
            )

            status_validation.add(
                "F2:F1000"
            )

            # -------------------------
            # Capacity validation
            # -------------------------

            capacity_validation = DataValidation(
                type="whole",
                operator="greaterThanOrEqual",
                formula1="0",
                allow_blank=False
            )

            capacity_validation.error = (
                "Capacity must be 0 or greater."
            )

            capacity_validation.errorTitle = (
                "Invalid Capacity"
            )

            sheet.add_data_validation(
                capacity_validation
            )

            capacity_validation.add(
                "D2:D1000"
            )

            # -------------------------
            # Freeze header
            # -------------------------

            sheet.freeze_panes = "A2"

            # -------------------------
            # Create response
            # -------------------------

            output = BytesIO()

            workbook.save(output)
            output.seek(0)

            application_logger.info(
                "room_excel_template_downloaded | "
                "school_id=%s",
                request.school.id
            )

            response = HttpResponse(
                output.getvalue(),
                content_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                )
            )

            response[
                "Content-Disposition"
            ] = (
                'attachment; '
                'filename="room_import_template.xlsx"'
            )

            return response

        except Exception:
            application_logger.exception(
                "room_excel_template_download_failed"
            )

            return CustomResponse.errorResponse(
                message="Failed to generate room Excel template"
            )

class TimetableCreateAPIView(APIView):
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = "timetable.create"

    def post(self, request):
        try:
            school = request.school

            name = request.data.get("name")
            academic_year_id = request.data.get("academic_year_id")
            term_id = request.data.get("term_id")
            branch_id = request.data.get("branch_id")
            status = request.data.get("status", "DRAFT")

            if not name:
                return CustomResponse.errorResponse(
                    message="Timetable name is required"
                )

            if not academic_year_id:
                return CustomResponse.errorResponse(
                    message="Academic year is required"
                )

            if not term_id:
                return CustomResponse.errorResponse(
                    message="Timetable term is required"
                )

            # -------------------------
            # UUID validation
            # -------------------------

            try:
                academic_year_uuid = uuid.UUID(
                    str(academic_year_id)
                )
                term_uuid = uuid.UUID(
                    str(term_id)
                )
            except (ValueError, AttributeError):
                return CustomResponse.errorResponse(
                    message="Invalid academic_year_id or term_id"
                )

            # -------------------------
            # Academic Year
            # -------------------------

            academic_year = AcademicYear.objects.filter(
                id=academic_year_uuid,
                school=school,
            ).first()

            if not academic_year:
                return CustomResponse.errorResponse(
                    message="Academic year not found"
                )

            # -------------------------
            # Timetable Term
            # -------------------------

            term = TimetableTerm.objects.filter(
                id=term_uuid,
                school=school,
                academic_year=academic_year,
            ).first()

            if not term:
                return CustomResponse.errorResponse(
                    message="Timetable term not found"
                )

            # -------------------------
            # Branch
            # -------------------------

            branch = None

            if branch_id:
                try:
                    branch_uuid = uuid.UUID(
                        str(branch_id)
                    )
                except (ValueError, AttributeError):
                    return CustomResponse.errorResponse(
                        message="Invalid branch_id"
                    )

                branch = Branch.objects.filter(
                    id=branch_uuid,
                    school=school,
                ).first()

                if not branch:
                    return CustomResponse.errorResponse(
                        message="Branch not found"
                    )

            # -------------------------
            # Status
            # -------------------------

            valid_statuses = dict(
                Timetable.status
            )

            if status not in valid_statuses:
                return CustomResponse.errorResponse(
                    message="Invalid timetable status"
                )

            # -------------------------
            # Duplicate
            # -------------------------

            if Timetable.objects.filter(
                school=school,
                academic_year=academic_year,
                term=term,
                name__iexact=name.strip(),
            ).exists():
                return CustomResponse.errorResponse(
                    message="Timetable with this name already exists"
                )

            # -------------------------
            # Create
            # -------------------------

            timetable = Timetable.objects.create(
                school=school,
                branch=branch,
                academic_year=academic_year,
                term=term,
                name=name.strip(),
                status=status,
            )

            application_logger.info(
                "timetable_created | timetable_id=%s | school_id=%s",
                timetable.id,
                school.id,
            )

            return CustomResponse.successResponse(
                message="Timetable created successfully",
                data={
                    "id": str(timetable.id),
                    "name": timetable.name,
                    "status": timetable.status,
                    "academic_year": {
                        "id": str(academic_year.id),
                        "name": academic_year.name,
                    },
                    "term": {
                        "id": str(term.id),
                        "name": term.name,
                    },
                    "branch": {
                        "id": str(branch.id),
                        "name": branch.name,
                    } if branch else None,
                },
            )

        except Exception:
            application_logger.exception(
                "timetable_create_failed"
            )

            return CustomResponse.errorResponse(
                message="Failed to create timetable"
            )


class TimetableListAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        HasPermission,
    ]

    required_permission = "timetable.view"

    def get(self, request):

        try:

            # ------------------------------------------------------
            # SCHOOL
            # ------------------------------------------------------

            school = request.school

            if not school:

                return CustomResponse.errorResponse(
                    description="School not found."
                )

            # ------------------------------------------------------
            # QUERY PARAMETERS
            # ------------------------------------------------------

            academic_year_id = (
                request.query_params.get(
                    "academic_year_id"
                )
            )

            term_id = (
                request.query_params.get(
                    "term_id"
                )
            )

            grade_id = (
                request.query_params.get(
                    "grade_id"
                )
            )

            section_id = (
                request.query_params.get(
                    "section_id"
                )
            )

            status = (
                request.query_params.get(
                    "status"
                )
            )

            # ------------------------------------------------------
            # GET TIMETABLES
            # ------------------------------------------------------

            timetables = (
                Timetable.objects
                .filter(
                    school=school,
                )
                .select_related(
                    "academic_year",
                    "term",
                    "branch",
                )
                .prefetch_related(
                    Prefetch(
                        "entries",
                        queryset=(
                            TimetableEntry.objects
                            .filter(
                                status=(
                                    TimetableEntry
                                    .Status
                                    .ACTIVE
                                )
                            )
                            .select_related(
                                "grade",
                                "section",
                                "subject",
                                "teacher",
                                "period",
                            )
                            .order_by(
                                "day_of_week",
                                "period__period_number",
                            )
                        ),
                    )
                )
                .order_by(
                    "-created_at"
                )
            )

            # ------------------------------------------------------
            # FILTER - ACADEMIC YEAR
            # ------------------------------------------------------

            if academic_year_id:

                timetables = timetables.filter(
                    academic_year_id=academic_year_id
                )

            # ------------------------------------------------------
            # FILTER - TERM
            # ------------------------------------------------------

            if term_id:

                timetables = timetables.filter(
                    term_id=term_id
                )

            # ------------------------------------------------------
            # FILTER - GRADE
            # ------------------------------------------------------

            if grade_id:

                timetables = timetables.filter(
                    entries__grade_id=grade_id
                )

            # ------------------------------------------------------
            # FILTER - SECTION
            # ------------------------------------------------------

            if section_id:

                timetables = timetables.filter(
                    entries__section_id=section_id
                )

            # ------------------------------------------------------
            # FILTER - STATUS
            # ------------------------------------------------------

            if status:

                timetables = timetables.filter(
                    status=status
                )

            # ------------------------------------------------------
            # REMOVE DUPLICATE TIMETABLES
            # ------------------------------------------------------

            timetables = timetables.distinct()

            # ------------------------------------------------------
            # PAGINATION
            # ------------------------------------------------------

            paginator = CustomPageNumberPagination()

            page = paginator.paginate_queryset(
                timetables,
                request,
            )

            # ------------------------------------------------------
            # RESPONSE DATA
            # ------------------------------------------------------

            data = []

            for timetable in page:

                entries = []

                for entry in timetable.entries.all():

                    entries.append({

                        "id": str(
                            entry.id
                        ),

                        "grade": {
                            "id": str(
                                entry.grade.id
                            ),
                            "name": (
                                entry.grade.name
                            ),
                        },

                        "section": {
                            "id": str(
                                entry.section.id
                            ),
                            "name": (
                                entry.section.name
                            ),
                        },

                        "day_of_week": (
                            entry.day_of_week
                        ),

                        "period": {

                            "id": str(
                                entry.period.id
                            ),

                            "period_number": (
                                entry.period
                                .period_number
                            ),

                            "name": getattr(
                                entry.period,
                                "name",
                                None,
                            ),

                            "start_time": getattr(
                                entry.period,
                                "start_time",
                                None,
                            ),

                            "end_time": getattr(
                                entry.period,
                                "end_time",
                                None,
                            ),

                            "is_break": getattr(
                                entry.period,
                                "is_break",
                                False,
                            ),
                        },

                        "subject": {
                            "id": str(
                                entry.subject.id
                            ),
                            "name": (
                                entry.subject.name
                            ),
                        },

                        "teacher": {
                            "id": str(
                                entry.teacher.id
                            ),
                            "name": (
                                entry.teacher.name
                            ),
                        },

                        "notes": entry.notes,

                        "status": (
                            entry.status
                        ),
                    })

                # --------------------------------------------------
                # TIMETABLE
                # --------------------------------------------------

                data.append({

                    "id": str(
                        timetable.id
                    ),

                    "name": (
                        timetable.name
                    ),

                    "academic_year": {

                        "id": str(
                            timetable
                            .academic_year
                            .id
                        ),

                        "name": (
                            timetable
                            .academic_year
                            .name
                        ),
                    },

                    "term": {

                        "id": str(
                            timetable
                            .term
                            .id
                        ),

                        "name": (
                            timetable
                            .term
                            .name
                        ),

                        "start_date": (
                            timetable
                            .term
                            .start_date
                        ),

                        "end_date": (
                            timetable
                            .term
                            .end_date
                        ),
                    },

                    "branch": (
                        {
                            "id": str(
                                timetable
                                .branch
                                .id
                            ),
                            "name": (
                                timetable
                                .branch
                                .name
                            ),
                        }
                        if timetable.branch
                        else None
                    ),

                    "status": (
                        timetable.status
                    ),

                    "entries": entries,

                    "created_at": (
                        timetable.created_at
                    ),

                    "updated_at": (
                        timetable.updated_at
                    ),
                })

            # ------------------------------------------------------
            # LOG
            # ------------------------------------------------------

            application_logger.info(
                "Timetables fetched successfully.",
                extra={

                    "school_id": str(
                        school.id
                    ),

                    "academic_year_id": (
                        academic_year_id
                    ),

                    "term_id": (
                        term_id
                    ),

                    "grade_id": (
                        grade_id
                    ),

                    "section_id": (
                        section_id
                    ),

                    "status": status,

                    "page": request
                    .query_params
                    .get(
                        "page",
                        1,
                    ),

                    "count": len(data),
                },
            )

            # ------------------------------------------------------
            # RESPONSE
            # ------------------------------------------------------

            return CustomResponse.successResponse(
                description=(
                    "Timetables fetched successfully."
                ),
                data=data,
            )

        except Exception as exc:

            application_logger.exception(
                "Error fetching timetables.",
                extra={

                    "school_id": (
                        str(
                            request.school.id
                        )
                        if getattr(
                            request,
                            "school",
                            None,
                        )
                        else None
                    ),

                    "academic_year_id": (
                        request.query_params.get(
                            "academic_year_id"
                        )
                    ),

                    "term_id": (
                        request.query_params.get(
                            "term_id"
                        )
                    ),

                    "grade_id": (
                        request.query_params.get(
                            "grade_id"
                        )
                    ),

                    "section_id": (
                        request.query_params.get(
                            "section_id"
                        )
                    ),

                    "error": str(exc),
                },
            )

            return CustomResponse.errorResponse(
                description=(
                    "Unable to fetch timetables."
                )
            )


class TimetableUpdateAPIView(APIView):
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = "timetable.update"

    @transaction.atomic
    def put(self, request, timetable_id):

        try:
            school = request.school

            if not school:
                return CustomResponse.errorResponse(
                    description="School not found."
                )

            # ----------------------------------------------------------
            # GET TIMETABLE
            # ----------------------------------------------------------

            timetable = (
                Timetable.objects
                .filter(
                    id=timetable_id,
                    school=school,
                )
                .first()
            )

            if not timetable:
                return CustomResponse.errorResponse(
                    description="Timetable not found."
                )

            # ----------------------------------------------------------
            # BASIC FIELDS
            # ----------------------------------------------------------

            if "status" in request.data:

                status = request.data.get("status")

                if status not in Timetable.Status.values:
                    return CustomResponse.errorResponse(
                        description="Invalid timetable status."
                    )

                timetable.status = status

            timetable.save()

            # ----------------------------------------------------------
            # ENTRIES
            # ----------------------------------------------------------

            entries = request.data.get("entries")

            if entries is not None:

                if not isinstance(entries, list):
                    return CustomResponse.errorResponse(
                        description="entries must be a list."
                    )

                # Remove existing entries
                TimetableEntry.objects.filter(
                    timetable=timetable
                ).update(
                    status=TimetableEntry.Status.INACTIVE
                )

                for entry_data in entries:

                    grade_id = entry_data.get("grade_id")
                    section_id = entry_data.get("section_id")
                    subject_id = entry_data.get("subject_id")
                    teacher_id = entry_data.get("teacher_id")
                    period_id = entry_data.get("period_id")
                    day_of_week = entry_data.get("day_of_week")
                    notes = entry_data.get("notes")

                    # --------------------------------------------------
                    # REQUIRED FIELDS
                    # --------------------------------------------------

                    required_fields = {
                        "grade_id": grade_id,
                        "section_id": section_id,
                        "subject_id": subject_id,
                        "teacher_id": teacher_id,
                        "period_id": period_id,
                        "day_of_week": day_of_week,
                    }

                    for field, value in required_fields.items():

                        if value in [None, ""]:
                            raise ValueError(
                                f"{field} is required."
                            )

                    # --------------------------------------------------
                    # VALIDATE DAY
                    # --------------------------------------------------

                    if day_of_week not in (
                        TimetableEntry.DayOfWeek.values
                    ):
                        raise ValueError(
                            "Invalid day_of_week."
                        )

                    # --------------------------------------------------
                    # GET GRADE
                    # --------------------------------------------------

                    grade = Grade.objects.filter(
                        id=grade_id,
                        school=school,
                    ).first()

                    if not grade:
                        raise ValueError(
                            "Grade not found."
                        )

                    # --------------------------------------------------
                    # GET SECTION
                    # --------------------------------------------------

                    section = Section.objects.filter(
                        id=section_id,
                        grade=grade,
                    ).first()

                    if not section:
                        raise ValueError(
                            "Section not found."
                        )

                    # --------------------------------------------------
                    # GET SUBJECT
                    # --------------------------------------------------

                    subject = Subject.objects.filter(
                        id=subject_id,
                        school=school,
                    ).first()

                    if not subject:
                        raise ValueError(
                            "Subject not found."
                        )

                    # --------------------------------------------------
                    # GET TEACHER
                    # --------------------------------------------------

                    teacher = Staff.objects.filter(
                        id=teacher_id,
                        school=school,
                    ).first()

                    if not teacher:
                        raise ValueError(
                            "Teacher not found."
                        )

                    # --------------------------------------------------
                    # GET PERIOD
                    # --------------------------------------------------

                    period = TimetablePeriod.objects.filter(
                        id=period_id,
                        school=school,
                    ).first()

                    if not period:
                        raise ValueError(
                            "Timetable period not found."
                        )

                    # --------------------------------------------------
                    # CREATE ENTRY
                    # --------------------------------------------------

                    TimetableEntry.objects.create(
                        timetable=timetable,
                        grade=grade,
                        section=section,
                        subject=subject,
                        teacher=teacher,
                        period=period,
                        day_of_week=day_of_week,
                        notes=notes,
                        status=TimetableEntry.Status.ACTIVE,
                    )

            application_logger.info(
                "Timetable updated successfully.",
                extra={
                    "school_id": str(school.id),
                    "timetable_id": str(timetable.id),
                    "performed_by": str(request.user.id),
                },
            )

            return CustomResponse.successResponse(
                description="Timetable updated successfully.",
                data={
                    "id": str(timetable.id),
                },
            )

        except IntegrityError as exc:

            application_logger.exception(
                "Timetable update integrity error.",
                extra={
                    "timetable_id": str(timetable_id),
                    "error": str(exc),
                },
            )

            return CustomResponse.errorResponse(
                description=(
                    "A timetable entry already exists for "
                    "this section, day and period."
                )
            )

        except Exception as exc:

            application_logger.exception(
                "Error updating timetable.",
                extra={
                    "timetable_id": str(timetable_id),
                    "error": str(exc),
                },
            )

            return CustomResponse.errorResponse(
                description=str(exc)
            )


class TimetableDeleteAPIView(APIView):
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = "timetable.delete"

    @transaction.atomic
    def delete(self, request, timetable_id):

        try:
            school = request.school

            if not school:
                return CustomResponse.errorResponse(
                    description="School not found."
                )

            timetable = (
                Timetable.objects
                .filter(
                    id=timetable_id,
                    school=school,
                )
                .first()
            )

            if not timetable:
                return CustomResponse.errorResponse(
                    description="Timetable not found."
                )

            # ----------------------------------------------------------
            # SOFT DELETE ENTRIES
            # ----------------------------------------------------------

            TimetableEntry.objects.filter(
                timetable=timetable
            ).update(
                status=TimetableEntry.Status.INACTIVE
            )

            # ----------------------------------------------------------
            # SOFT DELETE TIMETABLE
            # ----------------------------------------------------------

            timetable.soft_delete(request.user)

            application_logger.info(
                "Timetable deleted successfully.",
                extra={
                    "school_id": str(school.id),
                    "timetable_id": str(timetable.id),
                    "performed_by": str(request.user.id),
                },
            )

            return CustomResponse.successResponse(
                description="Timetable deleted successfully.",
                data={
                    "id": str(timetable.id),
                },
            )

        except Exception as exc:

            application_logger.exception(
                "Error deleting timetable.",
                extra={
                    "timetable_id": str(timetable_id),
                    "error": str(exc),
                },
            )

            return CustomResponse.errorResponse(
                description="Unable to delete timetable."
            )









class BulkUploadTimetableAPIView(APIView):

    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = "timetable.bulk_upload"

    parser_classes = [MultiPartParser]

    # --------------------------------------------------------------
    # TIME PARSER
    # --------------------------------------------------------------

    @staticmethod
    def parse_time(value):

        if value in [None, ""]:
            raise ValueError("Time is required.")

        # Excel time cell
        if isinstance(value, time):
            return value

        # Excel datetime cell
        if isinstance(value, datetime):
            return value.time()

        value = str(value).strip()

        for time_format in (
            "%H:%M",
            "%H:%M:%S",
            "%I:%M %p",
            "%I:%M:%S %p",
        ):
            try:
                return datetime.strptime(
                    value,
                    time_format,
                ).time()
            except ValueError:
                continue

        raise ValueError(
            f"Invalid time '{value}'. "
            f"Use HH:MM format."
        )

    # --------------------------------------------------------------
    # BOOLEAN PARSER
    # --------------------------------------------------------------

    @staticmethod
    def parse_boolean(value):

        if value in [None, ""]:
            return False

        if isinstance(value, bool):
            return value

        value = str(value).strip().lower()

        if value in (
            "yes",
            "true",
            "1",
            "y",
        ):
            return True

        if value in (
            "no",
            "false",
            "0",
            "n",
        ):
            return False

        raise ValueError(
            f"Invalid Is Break value '{value}'. "
            f"Use Yes or No."
        )

    # --------------------------------------------------------------
    # PERIOD NUMBER PARSER
    # --------------------------------------------------------------

    @staticmethod
    def parse_period_number(value):

        if value in [None, ""]:
            raise ValueError(
                "Period Number is required."
            )

        try:
            period_number = int(float(value))
        except (ValueError, TypeError):

            raise ValueError(
                f"Invalid period number '{value}'."
            )

        if period_number <= 0:
            raise ValueError(
                "Period Number must be greater than 0."
            )

        return period_number

    # --------------------------------------------------------------
    # POST
    # --------------------------------------------------------------

    def post(self, request):

        try:

            school = request.school

            application_logger.info(
                "timetable_bulk_upload_started",
                extra={
                    "performed_by": str(request.user.id),
                    "school_id": (
                        str(school.id)
                        if school
                        else None
                    ),
                },
            )

            # ----------------------------------------------------------
            # SCHOOL VALIDATION
            # ----------------------------------------------------------

            if not school:

                application_logger.warning(
                    "timetable_bulk_upload_failed",
                    extra={
                        "performed_by": str(
                            request.user.id
                        ),
                        "reason": "school_not_found",
                    },
                )

                return CustomResponse.errorResponse(
                    description="School not found."
                )

            # ----------------------------------------------------------
            # GET PARAMETERS
            # ----------------------------------------------------------

            academic_year_id = (
                request.query_params.get(
                    "academic_year_id"
                )
            )

            grade_id = request.query_params.get(
                "grade_id"
            )

            section_id = request.query_params.get(
                "section_id"
            )
            term_id = request.query_params.get("term_id")


            # ----------------------------------------------------------
            # PARAMETER VALIDATION
            # ----------------------------------------------------------

            if not academic_year_id:

                return CustomResponse.errorResponse(
                    description=(
                        "academic_year_id is required."
                    )
                )

            if not grade_id:

                return CustomResponse.errorResponse(
                    description="grade_id is required."
                )

            if not section_id:

                return CustomResponse.errorResponse(
                    description="section_id is required."
                )
            if not term_id:
                return CustomResponse.errorResponse(
                    description="term_id is required."
                )

            # ----------------------------------------------------------
            # UUID VALIDATION
            # ----------------------------------------------------------

            try:

                academic_year_uuid = uuid.UUID(
                    str(academic_year_id)
                )

            except (
                ValueError,
                TypeError,
                AttributeError,
            ):

                return CustomResponse.errorResponse(
                    description="Invalid academic_year_id."
                )

            try:

                grade_uuid = uuid.UUID(
                    str(grade_id)
                )

            except (
                ValueError,
                TypeError,
                AttributeError,
            ):

                return CustomResponse.errorResponse(
                    description="Invalid grade_id."
                )

            try:

                section_uuid = uuid.UUID(
                    str(section_id)
                )

            except (
                ValueError,
                TypeError,
                AttributeError,
            ):

                return CustomResponse.errorResponse(
                    description="Invalid section_id."
                )
            try:

                term_uuid = uuid.UUID(
                    str(term_id)
                )

            except (
                ValueError,
                TypeError,
                AttributeError,
            ):

                return CustomResponse.errorResponse(
                    description="Invalid term_id."
                )

            # ----------------------------------------------------------
            # ACADEMIC YEAR
            # ----------------------------------------------------------

            academic_year = AcademicYear.objects.filter(
                id=academic_year_uuid,
                school=school,
            ).first()

            if not academic_year:

                application_logger.warning(
                    "timetable_bulk_upload_failed",
                    extra={
                        "performed_by": str(
                            request.user.id
                        ),
                        "school_id": str(
                            school.id
                        ),
                        "academic_year_id": (
                            academic_year_id
                        ),
                        "reason": (
                            "academic_year_not_found"
                        ),
                    },
                )

                return CustomResponse.errorResponse(
                    description="Academic year not found."
                )

            # ----------------------------------------------------------
            # GRADE
            # ----------------------------------------------------------

            grade = Grade.objects.filter(
                id=grade_uuid,
                school=school,
            ).first()

            if not grade:

                application_logger.warning(
                    "timetable_bulk_upload_failed",
                    extra={
                        "performed_by": str(
                            request.user.id
                        ),
                        "school_id": str(
                            school.id
                        ),
                        "grade_id": str(
                            grade_id
                        ),
                        "reason": "grade_not_found",
                    },
                )

                return CustomResponse.errorResponse(
                    description="Grade not found."
                )

            # ----------------------------------------------------------
            # SECTION
            # ----------------------------------------------------------

            section = Section.objects.filter(
                id=section_uuid,
                grade=grade,
            ).first()

            if not section:

                application_logger.warning(
                    "timetable_bulk_upload_failed",
                    extra={
                        "performed_by": str(
                            request.user.id
                        ),
                        "school_id": str(
                            school.id
                        ),
                        "grade_id": str(
                            grade_id
                        ),
                        "section_id": str(
                            section_id
                        ),
                        "reason": (
                            "section_not_found"
                        ),
                    },
                )

                return CustomResponse.errorResponse(
                    description=(
                        "Section not found for "
                        "the selected grade."
                    )
                )
            term = TimetableTerm.objects.filter(
                id = term_uuid,
                school=school,
            ).first()
            if not term:
                application_logger.warning(
                    "timetable_bulk_upload_failed",
                    extra={
                        "performed_by": str(
                            request.user.id

                        ),
                        "school_id": str(
                            school.id),

                        "reason": (
                            "timetable term not found"
                        ),
                    }
                )
                return CustomResponse.errorResponse(
                    description="Timetable term not found."
                )

            # ----------------------------------------------------------
            # FILE
            # ----------------------------------------------------------

            file = request.FILES.get("file")

            if not file:

                application_logger.warning(
                    "timetable_bulk_upload_failed",
                    extra={
                        "performed_by": str(
                            request.user.id
                        ),
                        "school_id": str(
                            school.id
                        ),
                        "reason": (
                            "excel_file_missing"
                        ),
                    },
                )

                return CustomResponse.errorResponse(
                    description="Excel file is required."
                )

            # ----------------------------------------------------------
            # LOAD WORKBOOK
            # ----------------------------------------------------------

            try:

                workbook = load_workbook(
                    file,
                    data_only=True,
                )

            except Exception as exc:

                application_logger.exception(
                    "timetable_bulk_upload_failed",
                    extra={
                        "performed_by": str(
                            request.user.id
                        ),
                        "school_id": str(
                            school.id
                        ),
                        "file_name": file.name,
                        "reason": (
                            "invalid_excel_file"
                        ),
                        "error": str(exc),
                    },
                )

                return CustomResponse.errorResponse(
                    description="Invalid Excel file."
                )

            # ----------------------------------------------------------
            # PERIOD SHEET
            # ----------------------------------------------------------

            if "Periods" not in workbook.sheetnames:

                application_logger.warning(
                    "timetable_bulk_upload_failed",
                    extra={
                        "performed_by": str(
                            request.user.id
                        ),
                        "school_id": str(
                            school.id
                        ),
                        "file_name": file.name,
                        "reason": (
                            "period_sheet_missing"
                        ),
                    },
                )

                return CustomResponse.errorResponse(
                    description=(
                        "Periods sheet is required."
                    )
                )

            period_sheet = workbook["Periods"]

            # ----------------------------------------------------------
            # FIND TIMETABLE SHEET
            # ----------------------------------------------------------

            timetable_sheet = None

            for sheet_name in workbook.sheetnames:

                if sheet_name != "Periods":

                    timetable_sheet = workbook[
                        sheet_name
                    ]

                    break

            if not timetable_sheet:

                return CustomResponse.errorResponse(
                    description=(
                        "Timetable sheet is required."
                    )
                )

            # ==========================================================
            # PERIOD SHEET VALIDATION
            # ==========================================================

            period_headers = [
                str(cell.value).strip()
                if cell.value is not None
                else ""
                for cell in period_sheet[1]
            ]

            required_period_headers = [
                "Name",
                "Period Number",
                "Start Time",
                "End Time",
                "Is Break",
            ]

            for header in required_period_headers:

                if header not in period_headers:

                    application_logger.warning(
                        "timetable_bulk_upload_failed",
                        extra={
                            "performed_by": str(
                                request.user.id
                            ),
                            "school_id": str(
                                school.id
                            ),
                            "file_name": file.name,
                            "missing_header": header,
                            "reason": (
                                "period_column_missing"
                            ),
                        },
                    )

                    return CustomResponse.errorResponse(
                        description=(
                            f"{header} column is missing "
                            f"in Periods sheet."
                        )
                    )

            # ==========================================================
            # READ AND VALIDATE PERIODS
            # ==========================================================

            period_data = {}

            for row_number, row in enumerate(
                period_sheet.iter_rows(
                    min_row=2,
                    values_only=True,
                ),
                start=2,
            ):

                try:

                    data = dict(
                        zip(
                            period_headers,
                            row,
                        )
                    )

                    name = str(
                        data.get("Name") or ""
                    ).strip()

                    if not name:

                        raise ValueError(
                            "Name is required."
                        )

                    period_number = (
                        self.parse_period_number(
                            data.get(
                                "Period Number"
                            )
                        )
                    )

                    start_time = self.parse_time(
                        data.get("Start Time")
                    )

                    end_time = self.parse_time(
                        data.get("End Time")
                    )

                    if start_time >= end_time:

                        raise ValueError(
                            "Start Time must be "
                            "less than End Time."
                        )

                    is_break = self.parse_boolean(
                        data.get("Is Break")
                    )

                    if period_number in period_data:

                        raise ValueError(
                            f"Duplicate period number "
                            f"'{period_number}'."
                        )

                    period_data[period_number] = {
                        "name": name,
                        "period_number": (
                            period_number
                        ),
                        "start_time": start_time,
                        "end_time": end_time,
                        "is_break": is_break,
                    }

                except Exception as exc:

                    application_logger.warning(
                        "timetable_period_import_failed",
                        extra={
                            "performed_by": str(
                                request.user.id
                            ),
                            "school_id": str(
                                school.id
                            ),
                            "file_name": file.name,
                            "row_number": row_number,
                            "error": str(exc),
                        },
                    )

                    return CustomResponse.errorResponse(
                        description=(
                            f"Periods sheet row "
                            f"{row_number}: {str(exc)}"
                        )
                    )

            if not period_data:

                return CustomResponse.errorResponse(
                    description=(
                        "Periods sheet must contain "
                        "at least one period."
                    )
                )

            # ==========================================================
            # VALIDATE PERIOD OVERLAPPING
            # ==========================================================

            sorted_periods = sorted(
                period_data.values(),
                key=lambda item: item[
                    "start_time"
                ],
            )

            for index in range(
                len(sorted_periods) - 1
            ):

                current_period = (
                    sorted_periods[index]
                )

                next_period = (
                    sorted_periods[index + 1]
                )

                if (
                    current_period["end_time"]
                    > next_period["start_time"]
                ):

                    return CustomResponse.errorResponse(
                        description=(
                            "Periods overlap: "
                            f"Period "
                            f"{current_period['period_number']} "
                            f"and Period "
                            f"{next_period['period_number']}."
                        )
                    )

            # ==========================================================
            # TIMETABLE SHEET VALIDATION
            # ==========================================================

            timetable_headers = [
                str(cell.value).strip()
                if cell.value is not None
                else ""
                for cell in timetable_sheet[1]
            ]

            required_timetable_headers = [
                "Day",
                "Period",
                "Subject",
                "Teacher",
                "Notes",
            ]

            for header in required_timetable_headers:

                if header not in timetable_headers:

                    return CustomResponse.errorResponse(
                        description=(
                            f"{header} column is missing "
                            f"in timetable sheet."
                        )
                    )

            # ==========================================================
            # COUNTERS
            # ==========================================================

            period_created_count = 0
            period_existing_count = 0

            success_count = 0
            failed_count = 0

            errors = []

            # ==========================================================
            # CREATE PERIODS
            # ==========================================================

            try:

                with transaction.atomic():

                    periods = {}

                    for period_number, data in (
                        period_data.items()
                    ):

                        period = (
                            TimetablePeriod.objects.filter(
                                school=school,
                                period_number=(
                                    period_number
                                ),
                            ).first()
                        )

                        if period:

                            # Existing period must have
                            # the same timing.

                            if (
                                period.start_time
                                != data["start_time"]
                                or
                                period.end_time
                                != data["end_time"]
                            ):

                                raise ValueError(
                                    f"Period "
                                    f"'{period_number}' "
                                    f"already exists with "
                                    f"a different time."
                                )

                            periods[
                                period_number
                            ] = period

                            period_existing_count += 1

                        else:

                            period = (
                                TimetablePeriod.objects.create(
                                    school=school,
                                    name=data["name"],
                                    period_number=(
                                        period_number
                                    ),
                                    start_time=(
                                        data["start_time"]
                                    ),
                                    end_time=(
                                        data["end_time"]
                                    ),
                                    is_break=(
                                        data["is_break"]
                                    ),
                                    # status=(
                                    #     TimetablePeriod
                                    #     .Status
                                    #     .ACTIVE
                                    # ),
                                )
                            )

                            periods[
                                period_number
                            ] = period

                            period_created_count += 1

                    # ==================================================
                    # CREATE / GET TIMETABLE
                    # ==================================================

                    timetable, created = (
                        Timetable.objects.get_or_create(
                            school=school,
                            academic_year=(
                                academic_year
                            ),
                            term = term,
                            # grade=grade,
                            # section=section,
                            defaults={
                                "status": (
                                    Timetable
                                    .Status
                                    .ACTIVE
                                ),
                            },
                        )
                    )

                    # ==================================================
                    # PROCESS TIMETABLE ROWS
                    # ==================================================

                    for row_number, row in enumerate(
                        timetable_sheet.iter_rows(
                            min_row=2,
                            values_only=True,
                        ),
                        start=2,
                    ):

                        data = {}

                        try:

                            data = dict(
                                zip(
                                    timetable_headers,
                                    row,
                                )
                            )

                            # ------------------------------------------
                            # VALUES
                            # ------------------------------------------

                            day_of_week = str(
                                data.get("Day") or ""
                            ).strip().upper()

                            period_value = data.get(
                                "Period"
                            )

                            subject_name = str(
                                data.get("Subject")
                                or ""
                            ).strip()

                            teacher_name = str(
                                data.get("Teacher")
                                or ""
                            ).strip()

                            notes = data.get(
                                "Notes"
                            )

                            # ------------------------------------------
                            # REQUIRED VALUES
                            # ------------------------------------------

                            required_values = {
                                "Day": day_of_week,
                                "Period": period_value,
                                "Subject": subject_name,
                                "Teacher": teacher_name,
                            }

                            for field, value in (
                                required_values.items()
                            ):

                                if value in [
                                    None,
                                    "",
                                ]:

                                    raise ValueError(
                                        f"{field} is required."
                                    )

                            # ------------------------------------------
                            # DAY
                            # ------------------------------------------

                            if (
                                day_of_week
                                not in (
                                    TimetableEntry
                                    .DayOfWeek
                                    .values
                                )
                            ):

                                raise ValueError(
                                    f"Invalid day "
                                    f"'{day_of_week}'."
                                )

                            # ------------------------------------------
                            # PERIOD
                            # ------------------------------------------

                            try:

                                period_number = int(
                                    float(
                                        period_value
                                    )
                                )

                            except (
                                ValueError,
                                TypeError,
                            ):

                                raise ValueError(
                                    f"Invalid period "
                                    f"'{period_value}'."
                                )

                            period = periods.get(
                                period_number
                            )

                            if not period:

                                raise ValueError(
                                    f"Period "
                                    f"'{period_number}' "
                                    f"not found in "
                                    f"Periods sheet."
                                )

                            # ------------------------------------------
                            # SUBJECT
                            # ------------------------------------------

                            subject = (
                                Subject.objects.filter(
                                    school=school,
                                    name__iexact=(
                                        subject_name
                                    ),
                                ).first()
                            )

                            if not subject:

                                raise ValueError(
                                    f"Subject "
                                    f"'{subject_name}' "
                                    f"not found."
                                )

                            # ------------------------------------------
                            # TEACHER
                            # ------------------------------------------

                            teacher = (
                                Staff.objects.filter(
                                    school=school,
                                    name__iexact=(
                                        teacher_name
                                    ),
                                    status=(
                                        Staff.Status
                                        .ACTIVE
                                    ),
                                ).first()
                            )

                            if not teacher:

                                raise ValueError(
                                    f"Teacher "
                                    f"'{teacher_name}' "
                                    f"not found."
                                )

                            # ------------------------------------------
                            # DUPLICATE ENTRY
                            # ------------------------------------------

                            existing_entry = (
                                TimetableEntry.objects.filter(
                                    timetable=timetable,
                                    section=section,
                                    day_of_week=(
                                        day_of_week
                                    ),
                                    period=period,
                                ).first()
                            )

                            if existing_entry:

                                raise ValueError(
                                    "Timetable entry "
                                    "already exists for "
                                    f"{section.name}, "
                                    f"{day_of_week}, "
                                    f"Period "
                                    f"{period_number}."
                                )

                            # ------------------------------------------
                            # CREATE ENTRY
                            # ------------------------------------------

                            TimetableEntry.objects.create(
                                timetable=timetable,
                                grade=grade,
                                section=section,
                                subject=subject,
                                teacher=teacher,
                                period=period,
                                day_of_week=day_of_week,
                                notes=notes,
                                status=(
                                    TimetableEntry
                                    .Status
                                    .ACTIVE
                                ),
                            )

                            success_count += 1

                        except IntegrityError as exc:

                            failed_count += 1

                            errors.append({
                                "row": row_number,
                                "message": (
                                    "Duplicate timetable "
                                    "entry or invalid "
                                    "timetable data."
                                ),
                            })

                            application_logger.warning(
                                "timetable_bulk_upload_row_failed",
                                extra={
                                    "performed_by": str(
                                        request.user.id
                                    ),
                                    "school_id": str(
                                        school.id
                                    ),
                                    "file_name": (
                                        file.name
                                    ),
                                    "row_number": (
                                        row_number
                                    ),
                                    "error": str(exc),
                                },
                            )

                        except Exception as exc:

                            failed_count += 1

                            errors.append({
                                "row": row_number,
                                "message": str(exc),
                            })

                            application_logger.warning(
                                "timetable_bulk_upload_row_failed",
                                extra={
                                    "performed_by": str(
                                        request.user.id
                                    ),
                                    "school_id": str(
                                        school.id
                                    ),
                                    "file_name": (
                                        file.name
                                    ),
                                    "row_number": (
                                        row_number
                                    ),
                                    "day": data.get(
                                        "Day"
                                    ),
                                    "period": data.get(
                                        "Period"
                                    ),
                                    "subject": data.get(
                                        "Subject"
                                    ),
                                    "teacher": data.get(
                                        "Teacher"
                                    ),
                                    "error": str(exc),
                                },
                            )

            except ValueError as exc:

                application_logger.warning(
                    "timetable_bulk_upload_failed",
                    extra={
                        "performed_by": str(
                            request.user.id
                        ),
                        "school_id": str(
                            school.id
                        ),
                        "file_name": file.name,
                        "reason": "period_validation",
                        "error": str(exc),
                    },
                )

                return CustomResponse.errorResponse(
                    description=str(exc)
                )

            except IntegrityError as exc:

                application_logger.exception(
                    "timetable_bulk_upload_failed",
                    extra={
                        "performed_by": str(
                            request.user.id
                        ),
                        "school_id": str(
                            school.id
                        ),
                        "file_name": file.name,
                        "reason": "database_error",
                        "error": str(exc),
                    },
                )

                return CustomResponse.errorResponse(
                    description=(
                        "Unable to upload timetable."
                    )
                )

            except Exception as exc:

                application_logger.exception(
                    "timetable_bulk_upload_failed",
                    extra={
                        "performed_by": str(
                            request.user.id
                        ),
                        "school_id": str(
                            school.id
                        ),
                        "file_name": file.name,
                        "reason": "unexpected_error",
                        "error": str(exc),
                    },
                )

                return CustomResponse.errorResponse(
                    description=(
                        "Unable to upload timetable."
                    )
                )

            # ==========================================================
            # FINAL LOG
            # ==========================================================

            application_logger.info(
                "timetable_bulk_upload_completed",
                extra={
                    "performed_by": str(
                        request.user.id
                    ),
                    "school_id": str(
                        school.id
                    ),
                    "file_name": file.name,
                    "academic_year_id": str(
                        academic_year.id
                    ),
                    "grade_id": str(
                        grade.id
                    ),
                    "section_id": str(
                        section.id
                    ),
                    "period_created_count": (
                        period_created_count
                    ),
                    "period_existing_count": (
                        period_existing_count
                    ),
                    "success_count": success_count,
                    "failed_count": failed_count,
                },
            )

            # ==========================================================
            # RESPONSE
            # ==========================================================

            return CustomResponse.successResponse(
                description=(
                    "Timetable bulk upload completed."
                ),
                data={

                    "academic_year": {
                        "id": academic_year.id,
                        "name": academic_year.name,
                    },

                    "grade": {
                        "id": grade.id,
                        "name": grade.name,
                    },

                    "section": {
                        "id": section.id,
                        "name": section.name,
                    },

                    "periods": {
                        "created": (
                            period_created_count
                        ),
                        "existing": (
                            period_existing_count
                        ),
                        "total": (
                            period_created_count
                            + period_existing_count
                        ),
                    },

                    "timetable": {
                        "created": created,
                        "id": timetable.id,
                    },

                    "total_records": (
                        success_count
                        + failed_count
                    ),

                    "success_count": success_count,

                    "failed_count": failed_count,

                    "errors": errors,
                },
            )
        except Exception as exc:
            application_logger.exception(
                "timetable_create_exception",
                extra={
                    "requested_by": str(
                        request.user.id
                    ),
                    "error": str(exc),
                },
            )

            return CustomResponse.errorResponse(
                description=str(exc)
            )





# class DownloadTimetableTemplateAPIView(APIView):
#     permission_classes = [IsAuthenticated, HasPermission]
#     required_permission = "timetable.bulk_upload"
#
#     def get(self, request):
#
#         school = request.school
#
#         application_logger.info(
#             "timetable_template_download_started",
#             requested_by=str(request.user.id),
#             school_id=str(school.id) if school else None,
#         )
#
#         try:
#
#             # ------------------------------------------------------
#             # CREATE WORKBOOK
#             # ------------------------------------------------------
#
#             workbook = Workbook()
#
#             sheet = workbook.active
#             sheet.title = "Timetable"
#
#             # ------------------------------------------------------
#             # HEADERS
#             # ------------------------------------------------------
#
#             headers = [
#                 "Day",
#                 "Period",
#                 "Subject",
#                 "Teacher",
#                 "Notes",
#             ]
#
#             for column_number, header in enumerate(
#                 headers,
#                 start=1,
#             ):
#                 sheet.cell(
#                     row=1,
#                     column=column_number,
#                     value=header,
#                 )
#
#             # ------------------------------------------------------
#             # SAMPLE ROW
#             # ------------------------------------------------------
#
#             sample_row = [
#                 "MONDAY",
#                 1,
#                 "Mathematics",
#                 "John",
#                 "Regular class",
#             ]
#
#             for column_number, value in enumerate(
#                 sample_row,
#                 start=1,
#             ):
#                 sheet.cell(
#                     row=2,
#                     column=column_number,
#                     value=value,
#                 )
#
#             # ------------------------------------------------------
#             # SAVE EXCEL TO MEMORY
#             # ------------------------------------------------------
#
#             excel_file = BytesIO()
#
#             workbook.save(excel_file)
#
#             excel_file.seek(0)
#
#             # ------------------------------------------------------
#             # SAVE FILE TO STORAGE
#             # ------------------------------------------------------
#
#             file_name = (
#                 "templates/timetable_template.xlsx"
#             )
#
#             file_path = default_storage.save(
#                 file_name,
#                 ContentFile(
#                     excel_file.getvalue()
#                 ),
#             )
#
#             # ------------------------------------------------------
#             # FILE URL
#             # ------------------------------------------------------
#
#             file_url = (
#                 settings.MEDIA_URL
#                 + file_path
#             )
#
#         except Exception as e:
#
#             application_logger.exception(
#                 "timetable_template_download_failed",
#                 requested_by=str(request.user.id),
#                 school_id=(
#                     str(school.id)
#                     if school
#                     else None
#                 ),
#                 reason="template_generation_failed",
#                 error=str(e),
#             )
#
#             return CustomResponse.errorResponse(
#                 description=str(e),
#             )
#
#         # ----------------------------------------------------------
#         # SUCCESS LOG
#         # ----------------------------------------------------------
#
#         application_logger.info(
#             "timetable_template_generated",
#             requested_by=str(request.user.id),
#             school_id=str(school.id) if school else None,
#             file_path=file_path,
#             file_url=file_url,
#             total_columns=len(headers),
#         )
#
#         # ----------------------------------------------------------
#         # RESPONSE
#         # ----------------------------------------------------------
#
#         return CustomResponse.successResponse(
#             description=(
#                 "Timetable template generated successfully."
#             ),
#             data={
#                 "file_url": file_url,
#                 "file_path": file_path,
#             },
#         )



class DownloadTimetableTemplateAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        HasPermission,
    ]

    required_permission = "timetable.bulk_upload"

    def get(self, request):

        school = request.school

        application_logger.info(
            "timetable_template_download_started",
            extra={
                "requested_by": str(
                    request.user.id
                ),
                "school_id": (
                    str(school.id)
                    if school
                    else None
                ),
            },
        )

        try:

            # ------------------------------------------------------
            # SCHOOL VALIDATION
            # ------------------------------------------------------

            if not school:

                return CustomResponse.errorResponse(
                    description="School not found."
                )

            # ------------------------------------------------------
            # CREATE WORKBOOK
            # ------------------------------------------------------

            workbook = Workbook()

            # ======================================================
            # PERIODS SHEET
            # ======================================================

            period_sheet = workbook.active
            period_sheet.title = "Periods"

            period_headers = [
                "Name",
                "Period Number",
                "Start Time",
                "End Time",
                "Is Break",
            ]

            for column_number, header in enumerate(
                period_headers,
                start=1,
            ):

                period_sheet.cell(
                    row=1,
                    column=column_number,
                    value=header,
                )

            # ------------------------------------------------------
            # PERIOD SAMPLE DATA
            # ------------------------------------------------------

            period_sample_rows = [
                [
                    "Period 1",
                    1,
                    "09:00",
                    "09:45",
                    "No",
                ],
                [
                    "Period 2",
                    2,
                    "09:45",
                    "10:30",
                    "No",
                ],
                [
                    "Break",
                    3,
                    "10:30",
                    "10:45",
                    "Yes",
                ],
                [
                    "Period 4",
                    4,
                    "10:45",
                    "11:30",
                    "No",
                ],
            ]

            for row_number, row_data in enumerate(
                period_sample_rows,
                start=2,
            ):

                for column_number, value in enumerate(
                    row_data,
                    start=1,
                ):

                    period_sheet.cell(
                        row=row_number,
                        column=column_number,
                        value=value,
                    )

            # ------------------------------------------------------
            # PERIOD SHEET FORMATTING
            # ------------------------------------------------------

            period_sheet.freeze_panes = "A2"

            period_sheet.column_dimensions[
                "A"
            ].width = 20

            period_sheet.column_dimensions[
                "B"
            ].width = 18

            period_sheet.column_dimensions[
                "C"
            ].width = 15

            period_sheet.column_dimensions[
                "D"
            ].width = 15

            period_sheet.column_dimensions[
                "E"
            ].width = 15

            # ======================================================
            # TIMETABLE SHEET
            # ======================================================

            timetable_sheet = workbook.create_sheet(
                title="Timetable"
            )

            timetable_headers = [
                "Day",
                "Period",
                "Subject",
                "Teacher",
                "Notes",
            ]

            for column_number, header in enumerate(
                timetable_headers,
                start=1,
            ):

                timetable_sheet.cell(
                    row=1,
                    column=column_number,
                    value=header,
                )

            # ------------------------------------------------------
            # TIMETABLE SAMPLE DATA
            # ------------------------------------------------------

            timetable_sample_rows = [
                [
                    "MONDAY",
                    1,
                    "Mathematics",
                    "John",
                    "Regular class",
                ],
                [
                    "MONDAY",
                    2,
                    "English",
                    "David",
                    "",
                ],
                [
                    "TUESDAY",
                    1,
                    "Science",
                    "John",
                    "",
                ],
            ]

            for row_number, row_data in enumerate(
                timetable_sample_rows,
                start=2,
            ):

                for column_number, value in enumerate(
                    row_data,
                    start=1,
                ):

                    timetable_sheet.cell(
                        row=row_number,
                        column=column_number,
                        value=value,
                    )

            # ------------------------------------------------------
            # TIMETABLE SHEET FORMATTING
            # ------------------------------------------------------

            timetable_sheet.freeze_panes = "A2"

            timetable_sheet.column_dimensions[
                "A"
            ].width = 15

            timetable_sheet.column_dimensions[
                "B"
            ].width = 12

            timetable_sheet.column_dimensions[
                "C"
            ].width = 25

            timetable_sheet.column_dimensions[
                "D"
            ].width = 25

            timetable_sheet.column_dimensions[
                "E"
            ].width = 30

            # ------------------------------------------------------
            # SAVE EXCEL TO MEMORY
            # ------------------------------------------------------

            excel_file = BytesIO()

            workbook.save(excel_file)

            excel_file.seek(0)

            # ------------------------------------------------------
            # SAVE FILE TO STORAGE
            # ------------------------------------------------------

            file_name = (
                "templates/"
                "timetable_template.xlsx"
            )

            file_path = default_storage.save(
                file_name,
                ContentFile(
                    excel_file.getvalue()
                ),
            )

            # ------------------------------------------------------
            # FILE URL
            # ------------------------------------------------------

            file_url = (
                settings.MEDIA_URL
                + file_path
            )

        except Exception as exc:

            application_logger.exception(
                "timetable_template_download_failed",
                extra={
                    "requested_by": str(
                        request.user.id
                    ),
                    "school_id": (
                        str(school.id)
                        if school
                        else None
                    ),
                    "reason": (
                        "template_generation_failed"
                    ),
                    "error": str(exc),
                },
            )

            return CustomResponse.errorResponse(
                description=str(exc),
            )

        # ----------------------------------------------------------
        # SUCCESS LOG
        # ----------------------------------------------------------

        application_logger.info(
            "timetable_template_generated",
            extra={
                "requested_by": str(
                    request.user.id
                ),
                "school_id": str(
                    school.id
                ),
                "file_path": file_path,
                "file_url": file_url,
                "period_columns": len(
                    period_headers
                ),
                "timetable_columns": len(
                    timetable_headers
                ),
            },
        )

        # ----------------------------------------------------------
        # RESPONSE
        # ----------------------------------------------------------

        return CustomResponse.successResponse(
            description=(
                "Timetable template generated "
                "successfully."
            ),
            data={
                "file_url": file_url,
                "file_path": file_path,
                "sheets": [
                    "Periods",
                    "Timetable",
                ],
            },
        )


class TeacherScheduleListAPIView(APIView):
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = "teacher_schedule.view"

    def get(self, request):
        try:


            school = request.school

            application_logger.info(
                "teacher_schedule_list_started",
                extra={
                    "performed_by": str(request.user.id),
                    "school_id": str(school.id) if school else None,
                },
            )

            # ----------------------------------------------------------
            # SCHOOL VALIDATION
            # ----------------------------------------------------------

            if not school:
                return CustomResponse.errorResponse(
                    description="School not found."
                )

            # ----------------------------------------------------------
            # QUERY PARAMETERS
            # ----------------------------------------------------------

            search = request.GET.get("search")
            department = request.GET.get("department")
            status = request.GET.get("status")

            academic_year_id = request.GET.get("academic_year_id")
            grade_id = request.GET.get("grade_id")
            section_id = request.GET.get("section_id")
            branch_id = request.GET.get("branch_id")

            # ----------------------------------------------------------
            # BASE TEACHER QUERY
            # ----------------------------------------------------------

            teachers = Staff.objects.filter(
                school=school,
            )

            # ----------------------------------------------------------
            # SEARCH
            # ----------------------------------------------------------

            if search:

                teachers = teachers.filter(
                    Q(name__icontains=search)
                    | Q(employee_id__icontains=search)
                )

            # ----------------------------------------------------------
            # DEPARTMENT FILTER
            # ----------------------------------------------------------

            if department:

                valid_departments = {
                    value
                    for value, label in Staff.Department.choices
                }

                if department not in valid_departments:
                    return CustomResponse.errorResponse(
                        description="Invalid department."
                    )

                teachers = teachers.filter(
                    department=department
                )

            # ----------------------------------------------------------
            # TEACHER STATUS FILTER
            # ----------------------------------------------------------

            if status:
                teachers = teachers.filter(
                    status=status
                )

            # ----------------------------------------------------------
            # TIMETABLE ENTRY FILTERS
            # ----------------------------------------------------------

            entry_filters = {
                "status": TimetableEntry.Status.ACTIVE,
            }

            if academic_year_id:
                entry_filters[
                    "timetable__academic_year_id"
                ] = academic_year_id

            if grade_id:
                entry_filters[
                    "grade_id"
                ] = grade_id

            if section_id:
                entry_filters[
                    "section_id"
                ] = section_id

            if branch_id:
                entry_filters["timetable__branch_id"] = branch_id


            # ----------------------------------------------------------
            # PREFETCH TEACHER ENTRIES
            # ----------------------------------------------------------

            timetable_entries = (
                TimetableEntry.objects
                .filter(**entry_filters)
                .select_related(
                    "timetable",
                    "timetable__academic_year",
                    "timetable__branch",
                    "grade",
                    "section",
                    "subject",
                    "period",
                )
                .order_by(
                    "day_of_week",
                    "period__period_number",
                )
            )

            teachers = teachers.prefetch_related(
                Prefetch(
                    "timetable_entries",
                    queryset=timetable_entries,
                    to_attr="active_schedule_entries",
                )
            ).order_by("name")

            # ----------------------------------------------------------
            # PAGINATION
            # ----------------------------------------------------------

            paginator = CustomPageNumberPagination()

            page = paginator.paginate_queryset(
                teachers,
                request,
            )

            # ----------------------------------------------------------
            # RESPONSE
            # ----------------------------------------------------------

            data = []

            for teacher in page:

                entries = getattr(
                    teacher,
                    "active_schedule_entries",
                    [],
                )

                # ------------------------------------------------------
                # WEEKLY WORKLOAD
                # ------------------------------------------------------

                weekly_workload = len(entries)

                # ------------------------------------------------------
                # CLASS COUNT
                #
                # Distinct timetable / section combinations
                # ------------------------------------------------------

                class_count = len({
                    (
                        entry.timetable_id,
                        entry.section_id,
                    )
                    for entry in entries
                })

                # ------------------------------------------------------
                # SUBJECTS
                # ------------------------------------------------------

                subjects = {}

                for entry in entries:

                    subjects[entry.subject_id] = {
                        "id": entry.subject_id,
                        "name": entry.subject.name,
                    }

                # ------------------------------------------------------
                # SCHEDULE STATUS
                # ------------------------------------------------------

                if not entries:
                    schedule_status = "UNASSIGNED"

                elif (
                    hasattr(teacher, "status")
                    and teacher.status
                    != Staff.Status.ACTIVE
                ):
                    schedule_status = "DEACTIVATED"

                else:
                    schedule_status = "SYNCED"

                # ------------------------------------------------------
                # TEACHER RESPONSE
                # ------------------------------------------------------

                data.append({
                    "teacher": {
                        "id": teacher.id,
                        "employee_id": getattr(
                            teacher,
                            "employee_id",
                            None,
                        ),
                        "name": teacher.name,
                    },
                    "branch": (
                        {
                            "id": str(entries[0].timetable.branch.id),
                            "name": entries[0].timetable.branch.name,
                        }
                        if entries and entries[0].timetable.branch
                        else None
                    ),

                    "department": teacher.department if teacher.department else None,

                    "subjects": list(
                        subjects.values()
                    ),

                    "weekly_workload": {
                        "assigned_periods": weekly_workload,
                    },

                    "class_count": class_count,

                    "schedule_status": schedule_status,

                    "status": teacher.status,
                })


            return CustomResponse.successResponse(total=len(data),data=data,description="teachers data fetched successfully" )
        except Exception as exc:
            application_logger.exception(
                "teacher_schedule_exception",
                extra={
                    "requested_by": str(
                        request.user.id
                    ),
                    "error": str(exc),
                },
            )

            return CustomResponse.errorResponse(
                description=str(exc)
            )








class TeacherScheduleAPIView(APIView):
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = "teacher_schedule.view"

    def get(self, request, teacher_id):

        school = request.school

        application_logger.info(
            "teacher_schedule_view_started",
            extra={
                "performed_by": str(request.user.id),
                "school_id": str(school.id) if school else None,
                "teacher_id": str(teacher_id),
            },
        )

        # ----------------------------------------------------------
        # SCHOOL VALIDATION
        # ----------------------------------------------------------

        if not school:
            return CustomResponse.errorResponse(
                description="School not found."
            )

        # ----------------------------------------------------------
        # FILTERS
        # ----------------------------------------------------------

        academic_year_id = request.GET.get("academic_year_id")
        grade_id = request.GET.get("grade_id")
        section_id = request.GET.get("section_id")
        branch_id = request.GET.get("branch_id")
        department = request.GET.get("department")

        if not academic_year_id:
            return CustomResponse.errorResponse(
                description="academic_year_id is required."
            )

        # ----------------------------------------------------------
        # DEPARTMENT VALIDATION
        # ----------------------------------------------------------

        if department:

            valid_departments = {
                value
                for value, label in Staff.Department.choices
            }

            if department not in valid_departments:
                return CustomResponse.errorResponse(
                    description="Invalid department."
                )

        # ----------------------------------------------------------
        # TEACHER QUERY
        # ----------------------------------------------------------

        teacher_filters = {
            "id": teacher_id,
            "school": school,
        }

        if branch_id:
            teacher_filters["branch_id"] = branch_id

        if department:
            teacher_filters["department"] = department

        teacher = (
            Staff.objects
            .filter(**teacher_filters)
            .first()
        )

        if not teacher:
            return CustomResponse.errorResponse(
                description="Teacher not found."
            )

        # ----------------------------------------------------------
        # TIMETABLE ENTRY FILTERS
        # ----------------------------------------------------------

        entry_filters = {
            "timetable__school": school,
            "teacher": teacher,
            "timetable__academic_year_id": academic_year_id,
            "status": TimetableEntry.Status.ACTIVE,
        }

        if branch_id:
            entry_filters["timetable__branch_id"] = branch_id

        if grade_id:
            entry_filters["grade_id"] = grade_id

        if section_id:
            entry_filters["section_id"] = section_id

        # ----------------------------------------------------------
        # SCHEDULE QUERY
        # ----------------------------------------------------------

        entries = (
            TimetableEntry.objects
            .filter(**entry_filters)
            .select_related(
                "timetable",
                "timetable__academic_year",
                "timetable__branch",
                "grade",
                "section",
                "subject",
                "period",
            )
            .order_by(
                "day_of_week",
                "period__period_number",
            )
        )

        # ----------------------------------------------------------
        # GROUP BY DAY
        # ----------------------------------------------------------

        schedule = {
            day: []
            for day in TimetableEntry.DayOfWeek.values
        }

        for entry in entries:

            schedule[entry.day_of_week].append({

                "entry_id": str(entry.id),

                "day_of_week": entry.day_of_week,

                "period": {
                    "id": str(entry.period.id),
                    "period_number": entry.period.period_number,
                    "name": getattr(
                        entry.period,
                        "name",
                        None,
                    ),
                    "start_time": getattr(
                        entry.period,
                        "start_time",
                        None,
                    ),
                    "end_time": getattr(
                        entry.period,
                        "end_time",
                        None,
                    ),
                    "is_break": getattr(
                        entry.period,
                        "is_break",
                        False,
                    ),
                },

                "subject": {
                    "id": str(entry.subject.id),
                    "name": entry.subject.name,
                },

                "grade": {
                    "id": str(entry.grade.id),
                    "name": entry.grade.name,
                },

                "section": {
                    "id": str(entry.section.id),
                    "name": entry.section.name,
                },

                "notes": entry.notes,
            })

        # ----------------------------------------------------------
        # ACADEMIC YEAR
        # ----------------------------------------------------------

        academic_year = (
            entries[0].timetable.academic_year
            if entries
            else None
        )

        # ----------------------------------------------------------
        # BRANCH
        # ----------------------------------------------------------

        branch = (
            entries[0].timetable.branch
            if entries
            else teacher.branch
        )

        # ----------------------------------------------------------
        # RESPONSE
        # ----------------------------------------------------------

        return CustomResponse.successResponse(
            description="Teacher schedule fetched successfully.",
            data={

                "teacher": {
                    "id": str(teacher.id),
                    "employee_id": teacher.employee_id,
                    "name": teacher.name,
                },

                "department": (
                    {
                        "value": teacher.department,
                        "name": teacher.get_department_display(),
                    }
                    if teacher.department
                    else None
                ),

                "branch": (
                    {
                        "id": str(branch.id),
                        "name": branch.name,
                    }
                    if branch
                    else None
                ),

                "academic_year": {
                    "id": (
                        str(academic_year.id)
                        if academic_year
                        else academic_year_id
                    ),
                    "name": (
                        academic_year.name
                        if academic_year
                        else None
                    ),
                },

                "weekly_workload": len(entries),

                "class_count": len({
                    (
                        entry.grade_id,
                        entry.section_id,
                    )
                    for entry in entries
                }),

                "schedule": schedule,
            },
        )






class AssignTeacherSubstitutionAPIView(APIView):
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = "substitution.create"

    def post(self, request):

        school = request.school

        application_logger.info(
            "teacher_substitution_assignment_started",
            extra={
                "performed_by": str(request.user.id),
                "school_id": str(school.id) if school else None,
            },
        )

        # ==========================================================
        # SCHOOL VALIDATION
        # ==========================================================

        if not school:
            return CustomResponse.errorResponse(
                description="School not found."
            )

        # ==========================================================
        # REQUEST DATA
        # ==========================================================

        timetable_entry_id = request.data.get(
            "timetable_entry_id"
        )

        substitute_teacher_id = request.data.get(
            "substitute_teacher_id"
        )

        substitution_date = request.data.get(
            "date"
        )

        reason = request.data.get(
            "reason"
        )

        # ==========================================================
        # REQUIRED FIELD VALIDATION
        # ==========================================================

        if not timetable_entry_id:
            return CustomResponse.errorResponse(
                description="timetable_entry_id is required."
            )

        if not substitute_teacher_id:
            return CustomResponse.errorResponse(
                description="substitute_teacher_id is required."
            )

        if not substitution_date:
            return CustomResponse.errorResponse(
                description="date is required."
            )

        # ==========================================================
        # UUID VALIDATION
        # ==========================================================

        try:
            timetable_entry_uuid = uuid.UUID(
                str(timetable_entry_id)
            )
        except (ValueError, TypeError, AttributeError):

            return CustomResponse.errorResponse(
                description="Invalid timetable_entry_id."
            )

        try:
            substitute_teacher_uuid = uuid.UUID(
                str(substitute_teacher_id)
            )
        except (ValueError, TypeError, AttributeError):

            return CustomResponse.errorResponse(
                description="Invalid substitute_teacher_id."
            )

        # ==========================================================
        # DATE VALIDATION
        # ==========================================================

        try:
            substitution_date = date.fromisoformat(
                str(substitution_date)
            )
        except (ValueError, TypeError):

            return CustomResponse.errorResponse(
                description=(
                    "Invalid date. "
                    "Use YYYY-MM-DD format."
                )
            )

        # ==========================================================
        # GET TIMETABLE ENTRY
        # ==========================================================

        timetable_entry = (
            TimetableEntry.objects
            .filter(
                id=timetable_entry_uuid,
                timetable__school=school,
                status=TimetableEntry.Status.ACTIVE,
            )
            .select_related(
                "timetable",
                "timetable__academic_year",
                "grade",
                "section",
                "subject",
                "teacher",
                "period",
            )
            .first()
        )

        if not timetable_entry:

            application_logger.warning(
                "teacher_substitution_assignment_failed",
                extra={
                    "performed_by": str(request.user.id),
                    "school_id": str(school.id),
                    "timetable_entry_id": str(
                        timetable_entry_id
                    ),
                    "reason": "timetable_entry_not_found",
                },
            )

            return CustomResponse.errorResponse(
                description="Timetable entry not found."
            )

        # ==========================================================
        # CHECK DATE AGAINST TIMETABLE DAY
        # ==========================================================

        selected_day = (
            substitution_date
            .strftime("%A")
            .upper()
        )

        if selected_day != timetable_entry.day_of_week:

            return CustomResponse.errorResponse(
                description=(
                    f"Selected date is {selected_day}, "
                    f"but the timetable entry is scheduled "
                    f"for {timetable_entry.day_of_week}."
                )
            )

        # ==========================================================
        # ORIGINAL TEACHER
        # ==========================================================

        original_teacher = timetable_entry.teacher

        # ==========================================================
        # GET SUBSTITUTE TEACHER
        # ==========================================================

        substitute_teacher = (
            Staff.objects
            .filter(
                id=substitute_teacher_uuid,
                school=school,
                status=Staff.Status.ACTIVE,
            )
            .first()
        )

        if not substitute_teacher:

            return CustomResponse.errorResponse(
                description=(
                    "Substitute teacher not found "
                    "or inactive."
                )
            )

        # ==========================================================
        # SAME TEACHER VALIDATION
        # ==========================================================

        if (
            original_teacher.id
            == substitute_teacher.id
        ):

            return CustomResponse.errorResponse(
                description=(
                    "Substitute teacher cannot be "
                    "the original teacher."
                )
            )

        # ==========================================================
        # EXISTING SUBSTITUTION FOR THIS CLASS / DATE
        # ==========================================================

        existing_substitution = (
            TimetableSubstitution.objects
            .filter(
                timetable_entry=timetable_entry,
                date=substitution_date,
                status=(
                    TimetableSubstitution
                    .Status
                    .SCHEDULED
                ),
            )
            .first()
        )

        if existing_substitution:

            return CustomResponse.errorResponse(
                description=(
                    "A substitution is already scheduled "
                    "for this class on the selected date."
                )
            )

        # ==========================================================
        # SUBSTITUTE TEACHER REGULAR TIMETABLE CONFLICT
        # ==========================================================

        timetable_conflict = (
            TimetableEntry.objects
            .filter(
                timetable__school=school,
                timetable__academic_year_id=(
                    timetable_entry
                    .timetable
                    .academic_year_id
                ),
                teacher=substitute_teacher,
                day_of_week=timetable_entry.day_of_week,
                period=timetable_entry.period,
                status=TimetableEntry.Status.ACTIVE,
            )
            .exclude(
                id=timetable_entry.id
            )
            .select_related(
                "grade",
                "section",
                "subject",
            )
            .first()
        )

        if timetable_conflict:

            application_logger.warning(
                "teacher_substitution_period_conflict",
                extra={
                    "performed_by": str(request.user.id),
                    "school_id": str(school.id),
                    "substitute_teacher_id": str(
                        substitute_teacher.id
                    ),
                    "timetable_entry_id": str(
                        timetable_entry.id
                    ),
                    "conflicting_entry_id": str(
                        timetable_conflict.id
                    ),
                },
            )

            return CustomResponse.errorResponse(
                description=(
                    f"{substitute_teacher.name} is already "
                    f"scheduled for "
                    f"{timetable_conflict.subject.name} "
                    f"({timetable_conflict.grade.name} "
                    f"{timetable_conflict.section.name}) "
                    f"during Period "
                    f"{timetable_entry.period.period_number}."
                )
            )

        # ==========================================================
        # SUBSTITUTE TEACHER EXISTING SUBSTITUTION CONFLICT
        # ==========================================================

        substitution_conflict = (
            TimetableSubstitution.objects
            .filter(
                substitute_teacher=substitute_teacher,
                date=substitution_date,
                status=(
                    TimetableSubstitution
                    .Status
                    .SCHEDULED
                ),
                timetable_entry__period=(
                    timetable_entry.period
                ),
                timetable_entry__day_of_week=(
                    timetable_entry.day_of_week
                ),
            )
            .select_related(
                "timetable_entry",
                "timetable_entry__subject",
                "timetable_entry__grade",
                "timetable_entry__section",
            )
            .first()
        )

        if substitution_conflict:

            conflict_entry = (
                substitution_conflict
                .timetable_entry
            )

            return CustomResponse.errorResponse(
                description=(
                    f"{substitute_teacher.name} is already "
                    f"assigned as a substitute for "
                    f"{conflict_entry.subject.name} "
                    f"({conflict_entry.grade.name} "
                    f"{conflict_entry.section.name}) "
                    f"during Period "
                    f"{timetable_entry.period.period_number}."
                )
            )

        # ==========================================================
        # WEEKLY WORKLOAD
        # ==========================================================

        # Current regular timetable periods for the teacher
        current_regular_periods = (
            TimetableEntry.objects
            .filter(
                timetable__school=school,
                timetable__academic_year_id=(
                    timetable_entry
                    .timetable
                    .academic_year_id
                ),
                teacher=substitute_teacher,
                status=TimetableEntry.Status.ACTIVE,
            )
            .count()
        )

        # Existing scheduled substitutions for the teacher
        current_substitution_periods = (
            TimetableSubstitution.objects
            .filter(
                substitute_teacher=substitute_teacher,
                status=(
                    TimetableSubstitution
                    .Status
                    .SCHEDULED
                ),
                timetable_entry__timetable__academic_year_id=(
                    timetable_entry
                    .timetable
                    .academic_year_id
                ),
            )
            .count()
        )

        current_workload = (
            current_regular_periods
            + current_substitution_periods
        )

        workload_after_substitution = (
            current_workload + 1
        )

        # ----------------------------------------------------------
        # MAXIMUM WORKLOAD
        #
        # There is no maximum workload field in the models
        # provided so far. Therefore we don't reject the
        # substitution based on a hardcoded maximum.
        # ----------------------------------------------------------

        maximum_workload = None

        # ==========================================================
        # VALIDATION RESULT
        # ==========================================================

        validation = {
            "period_conflict": False,
            "workload_limit_exceeded": False,
            "subject_expertise_match": None,
        }

        # ==========================================================
        # CREATE SUBSTITUTION
        # ==========================================================

        try:

            with transaction.atomic():

                substitution = (
                    TimetableSubstitution.objects.create(
                        timetable_entry=timetable_entry,
                        date=substitution_date,
                        original_teacher=original_teacher,
                        substitute_teacher=substitute_teacher,
                        reason=reason,
                        status=(
                            TimetableSubstitution
                            .Status
                            .SCHEDULED
                        ),
                        created_by=request.user,
                    )
                )

        except IntegrityError:

            application_logger.warning(
                "teacher_substitution_duplicate",
                extra={
                    "performed_by": str(request.user.id),
                    "school_id": str(school.id),
                    "timetable_entry_id": str(
                        timetable_entry.id
                    ),
                    "date": str(
                        substitution_date
                    ),
                },
            )

            return CustomResponse.errorResponse(
                description=(
                    "A substitution already exists "
                    "for this timetable entry on "
                    "the selected date."
                )
            )

        except Exception as exc:

            application_logger.exception(
                "teacher_substitution_assignment_failed",
                extra={
                    "performed_by": str(request.user.id),
                    "school_id": str(school.id),
                    "timetable_entry_id": str(
                        timetable_entry.id
                    ),
                    "substitute_teacher_id": str(
                        substitute_teacher.id
                    ),
                    "error": str(exc),
                },
            )

            return CustomResponse.errorResponse(
                description=(
                    "Unable to assign teacher substitution."
                )
            )

        # ==========================================================
        # RESPONSE DATA
        # ==========================================================

        data = {
            "id": str(substitution.id),

            "date": substitution.date,

            # ------------------------------------------------------
            # ABSENT TEACHER
            # ------------------------------------------------------

            "absent_teacher": {
                "id": str(
                    original_teacher.id
                ),
                "employee_id": getattr(
                    original_teacher,
                    "employee_id",
                    None,
                ),
                "name": original_teacher.name,
            },

            # ------------------------------------------------------
            # SUBSTITUTE TEACHER
            # ------------------------------------------------------

            "substitute_teacher": {
                "id": str(
                    substitute_teacher.id
                ),
                "employee_id": getattr(
                    substitute_teacher,
                    "employee_id",
                    None,
                ),
                "name": substitute_teacher.name,
            },

            # ------------------------------------------------------
            # ASSIGNMENT
            # ------------------------------------------------------

            "assignment": {
                "grade": {
                    "id": str(
                        timetable_entry
                        .grade
                        .id
                    ),
                    "name": (
                        timetable_entry
                        .grade
                        .name
                    ),
                },

                "section": {
                    "id": str(
                        timetable_entry
                        .section
                        .id
                    ),
                    "name": (
                        timetable_entry
                        .section
                        .name
                    ),
                },

                "subject": {
                    "id": str(
                        timetable_entry
                        .subject
                        .id
                    ),
                    "name": (
                        timetable_entry
                        .subject
                        .name
                    ),
                },

                "day_of_week": (
                    timetable_entry
                    .day_of_week
                ),

                "period": {
                    "id": str(
                        timetable_entry
                        .period
                        .id
                    ),
                    "period_number": (
                        timetable_entry
                        .period
                        .period_number
                    ),
                    "start_time": getattr(
                        timetable_entry.period,
                        "start_time",
                        None,
                    ),
                    "end_time": getattr(
                        timetable_entry.period,
                        "end_time",
                        None,
                    ),
                },
            },

            # ------------------------------------------------------
            # WORKLOAD
            # ------------------------------------------------------

            "workload": {
                "current_periods": current_workload,
                "after_substitution": (
                    workload_after_substitution
                ),
                "maximum_periods": maximum_workload,
            },

            # ------------------------------------------------------
            # VALIDATION
            # ------------------------------------------------------

            "validation": validation,

            "reason": substitution.reason,

            "status": substitution.status,
        }

        # ==========================================================
        # LOG
        # ==========================================================

        application_logger.info(
            "teacher_substitution_assignment_completed",
            extra={
                "performed_by": str(request.user.id),
                "school_id": str(school.id),
                "substitution_id": str(
                    substitution.id
                ),
                "timetable_entry_id": str(
                    timetable_entry.id
                ),
                "original_teacher_id": str(
                    original_teacher.id
                ),
                "substitute_teacher_id": str(
                    substitute_teacher.id
                ),
                "subject_id": str(
                    timetable_entry.subject_id
                ),
                "date": str(
                    substitution_date
                ),
            },
        )

        # ==========================================================
        # RESPONSE
        # ==========================================================

        return CustomResponse.successResponse(
            data=data,
            description=(
                "Teacher substitution assigned successfully."
            ),
        )


class TeacherSubstitutionListAPIView(APIView):
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = "substitution.view"

    def get(self, request):

        school = request.school

        application_logger.info(
            "teacher_substitution_list_started",
            extra={
                "performed_by": str(request.user.id),
                "school_id": str(school.id) if school else None,
            },
        )

        # ==========================================================
        # SCHOOL VALIDATION
        # ==========================================================

        if not school:
            return CustomResponse.errorResponse(
                description="School not found."
            )

        # ==========================================================
        # QUERY PARAMETERS
        # ==========================================================

        date_value = request.GET.get("date")
        status = request.GET.get("status")
        teacher_id = request.GET.get("teacher_id")
        grade_id = request.GET.get("grade_id")
        section_id = request.GET.get("section_id")
        search = request.GET.get("search")
        academic_year_id = request.GET.get("academic_year_id")

        # ==========================================================
        # BASE QUERY
        # ==========================================================

        substitutions = (
            TimetableSubstitution.objects
            .filter(
                timetable_entry__timetable__school=school,
            )
            .select_related(
                "timetable_entry",
                "timetable_entry__timetable",
                "timetable_entry__timetable__academic_year",
                "timetable_entry__grade",
                "timetable_entry__section",
                "timetable_entry__subject",
                "timetable_entry__period",
                "original_teacher",
                "substitute_teacher",
            )
        )

        # ==========================================================
        # ACADEMIC YEAR FILTER
        # ==========================================================

        if academic_year_id:
            substitutions = substitutions.filter(
                timetable_entry__timetable__academic_year_id=(
                    academic_year_id
                )
            )

        # ==========================================================
        # DATE FILTER
        # ==========================================================

        if date_value:
            substitutions = substitutions.filter(
                date=date_value
            )

        # ==========================================================
        # STATUS FILTER
        # ==========================================================

        if status:
            substitutions = substitutions.filter(
                status=status
            )

        # ==========================================================
        # TEACHER FILTER
        #
        # Matches either original teacher or substitute teacher.
        # ==========================================================

        if teacher_id:

            substitutions = substitutions.filter(
                Q(original_teacher_id=teacher_id)
                | Q(substitute_teacher_id=teacher_id)
            )

        # ==========================================================
        # GRADE FILTER
        # ==========================================================

        if grade_id:
            substitutions = substitutions.filter(
                timetable_entry__grade_id=grade_id
            )

        # ==========================================================
        # SECTION FILTER
        # ==========================================================

        if section_id:
            substitutions = substitutions.filter(
                timetable_entry__section_id=section_id
            )

        # ==========================================================
        # SEARCH
        #
        # Search by original teacher, substitute teacher,
        # subject, grade, section.
        # ==========================================================

        if search:

            substitutions = substitutions.filter(
                Q(original_teacher__name__icontains=search)
                | Q(
                    substitute_teacher__name__icontains=search
                )
                | Q(
                    timetable_entry__subject__name__icontains=search
                )
                | Q(
                    timetable_entry__grade__name__icontains=search
                )
                | Q(
                    timetable_entry__section__name__icontains=search
                )
            )

        # ==========================================================
        # ORDERING
        # ==========================================================

        substitutions = substitutions.order_by(
            "-date",
            "timetable_entry__period__period_number",
        )

        # ==========================================================
        # PAGINATION
        # ==========================================================

        paginator = CustomPageNumberPagination()

        page = paginator.paginate_queryset(
            substitutions,
            request,
        )

        # ==========================================================
        # RESPONSE DATA
        # ==========================================================

        data = []

        for substitution in page:

            entry = substitution.timetable_entry

            data.append({
                "id": str(
                    substitution.id
                ),

                "date": substitution.date,

                "status": substitution.status,

                # --------------------------------------------------
                # ABSENT / ORIGINAL TEACHER
                # --------------------------------------------------

                "absent_teacher": {
                    "id": str(
                        substitution
                        .original_teacher
                        .id
                    ),
                    "employee_id": getattr(
                        substitution.original_teacher,
                        "employee_id",
                        None,
                    ),
                    "name": (
                        substitution
                        .original_teacher
                        .name
                    ),
                },

                # --------------------------------------------------
                # SUBSTITUTE TEACHER
                # --------------------------------------------------

                "substitute_teacher": {
                    "id": str(
                        substitution
                        .substitute_teacher
                        .id
                    ),
                    "employee_id": getattr(
                        substitution.substitute_teacher,
                        "employee_id",
                        None,
                    ),
                    "name": (
                        substitution
                        .substitute_teacher
                        .name
                    ),
                },

                # --------------------------------------------------
                # CLASS
                # --------------------------------------------------

                "class": {
                    "grade": {
                        "id": str(
                            entry.grade.id
                        ),
                        "name": (
                            entry.grade.name
                        ),
                    },
                    "section": {
                        "id": str(
                            entry.section.id
                        ),
                        "name": (
                            entry.section.name
                        ),
                    },
                },

                # --------------------------------------------------
                # SUBJECT
                # --------------------------------------------------

                "subject": {
                    "id": str(
                        entry.subject.id
                    ),
                    "name": (
                        entry.subject.name
                    ),
                },

                # --------------------------------------------------
                # PERIOD
                # --------------------------------------------------

                "period": {
                    "id": str(
                        entry.period.id
                    ),
                    "period_number": (
                        entry.period
                        .period_number
                    ),
                    "start_time": getattr(
                        entry.period,
                        "start_time",
                        None,
                    ),
                    "end_time": getattr(
                        entry.period,
                        "end_time",
                        None,
                    ),
                },

                "day_of_week": (
                    entry.day_of_week
                ),

                "reason": substitution.reason,

                # --------------------------------------------------
                # ACADEMIC YEAR
                # --------------------------------------------------

                "academic_year": {
                    "id": str(
                        entry
                        .timetable
                        .academic_year
                        .id
                    ),
                    "name": (
                        entry
                        .timetable
                        .academic_year
                        .name
                    ),
                },
            })

        # ==========================================================
        # PAGINATION
        # ==========================================================

        # pagination = {
        #     "count": paginator.page.paginator.count,
        #     "page": paginator.page.number,
        #     "page_size": (
        #         paginator.page
        #         .paginator
        #         .per_page
        #     ),
        #     "total_pages": (
        #         paginator.page
        #         .paginator
        #         .num_pages
        #     ),
        # }

        # ==========================================================
        # LOG
        # ==========================================================

        application_logger.info(
            "teacher_substitution_list_completed",
            extra={
                "performed_by": str(request.user.id),
                "school_id": str(school.id),
                "count": len(data),
            },
        )

        # ==========================================================
        # RESPONSE
        # ==========================================================

        return CustomResponse.successResponse(
            data= data,

            description=(
                "Teacher substitutions fetched successfully."
            ),
        )
